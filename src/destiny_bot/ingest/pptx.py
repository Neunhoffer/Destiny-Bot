"""PowerPoint presentations (``.pptx``).

Each slide becomes one section, "Slide N: <title>". Its blocks follow reading
order (top to bottom, then left to right, looking inside groups): the title,
text boxes (bulleted lines become list items), tables, pictures with alt text,
chart titles and SmartArt text. Speaker notes come last, marked with
``meta["speaker_notes"]``.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree

from ..models import Block, BlockKind, Document
from . import IngestError, ProgressFn
from .common import clean_display_text, doc_id_for_file, finalize_sections, make_section, title_from_filename
from .text import table_block, text_block

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"
_XML_PARSER = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=True)
_DEFAULT_TITLES = frozenset({"powerpoint presentation", "presentation", "slide 1"})


def _paragraph_text(p) -> str:
    """Text of a DrawingML paragraph (runs, fields, line breaks)."""
    parts: list[str] = []
    for child in p.iterchildren():
        if child.tag in (_A + "r", _A + "fld"):
            t = child.find(_A + "t")
            if t is not None and t.text:
                parts.append(t.text)
        elif child.tag == _A + "br":
            parts.append("\n")
    return "".join(parts).replace("\v", "\n")


def _bullet_flag(ppr) -> bool | None:
    if ppr is None:
        return None
    if ppr.find(_A + "buNone") is not None:
        return False
    if ppr.find(_A + "buChar") is not None or ppr.find(_A + "buAutoNum") is not None or ppr.find(_A + "buBlip") is not None:
        return True
    return None


def _alt_text(shape) -> str:
    c_nv_pr = shape._element.find(".//" + _P + "cNvPr")
    if c_nv_pr is None:
        return ""
    return (c_nv_pr.get("descr") or "").strip()


def _page_list(numbers: list[int]) -> str:
    """[3, 4, 5, 9] -> "3-5, 9"."""
    ranges: list[str] = []
    start = prev = None
    for n in numbers + [None]:  # type: ignore[list-item]
        if start is None:
            start = prev = n
        elif n is not None and n == prev + 1:
            prev = n
        else:
            ranges.append(str(start) if start == prev else f"{start}-{prev}")
            start = prev = n
    return ", ".join(ranges)


class _SlideReader:
    def __init__(self, presentation, slide, number: int) -> None:
        self.presentation = presentation
        self.slide = slide
        self.number = number
        self.blocks: list[Block] = []
        height = presentation.slide_height or 6858000
        self.tolerance = int(height * 0.04)  # shapes this close vertically count as one row

    def add(self, block: Block | None) -> None:
        if block is not None:
            block.page = self.number
            self.blocks.append(block)

    # ------------------------------------------------------------ ordering

    def ordered(self, shapes) -> list:
        items = sorted(((s.top or 0, s.left or 0, i, s) for i, s in enumerate(shapes)), key=lambda t: (t[0], t[1], t[2]))
        rows: list[list[tuple]] = []
        for item in items:
            if rows and abs(item[0] - rows[-1][0][0]) <= self.tolerance:
                rows[-1].append(item)
            else:
                rows.append([item])
        return [item[3] for row in rows for item in sorted(row, key=lambda t: (t[1], t[2]))]

    # ------------------------------------------------------------ shapes

    @staticmethod
    def placeholder_type(shape):
        if not getattr(shape, "is_placeholder", False):
            return None
        try:
            return shape.placeholder_format.type
        except Exception:
            return None

    def shapes(self, shapes, title_shape=None) -> None:
        from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
        from pptx.shapes.picture import Picture

        skip_types = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.HEADER}
        for shape in self.ordered(list(shapes)):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            ph_type = self.placeholder_type(shape)
            if ph_type in skip_types:
                continue
            try:
                shape_type = shape.shape_type
            except Exception:
                shape_type = None
            if shape_type == MSO_SHAPE_TYPE.GROUP:
                self.shapes(shape.shapes)
                continue
            if getattr(shape, "has_table", False):
                self.table(shape)
                continue
            if getattr(shape, "has_chart", False):
                self.chart(shape)
                continue
            if isinstance(shape, Picture) or shape_type == MSO_SHAPE_TYPE.PICTURE:
                self.add(text_block(BlockKind.FIGURE, _alt_text(shape)))
                continue
            if self.smartart(shape):
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                self.text_frame(shape, ph_type)
            elif _alt_text(shape) and shape_type in (MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT):
                self.add(text_block(BlockKind.FIGURE, _alt_text(shape)))

    def text_frame(self, shape, ph_type) -> None:
        from pptx.enum.shapes import PP_PLACEHOLDER

        bulleted_by_default = ph_type in (
            PP_PLACEHOLDER.BODY,
            PP_PLACEHOLDER.OBJECT,
            PP_PLACEHOLDER.VERTICAL_BODY,
            PP_PLACEHOLDER.VERTICAL_OBJECT,
        )
        for paragraph in shape.text_frame.paragraphs:
            text = _paragraph_text(paragraph._p)
            if not text.strip():
                continue
            level = paragraph.level or 0
            bullet = _bullet_flag(paragraph._p.find(_A + "pPr"))
            if bullet is None and bulleted_by_default:
                inherited = self.inherited_bullet(shape, level)
                bullet = True if inherited is None else inherited
            if level > 0 or bullet:
                self.add(text_block(BlockKind.LIST_ITEM, text, meta={"list_level": level + 1}))
            else:
                self.add(text_block(BlockKind.PARAGRAPH, text))

    def inherited_bullet(self, shape, level: int) -> bool | None:
        """Bullet setting a body placeholder inherits from its layout / master."""
        tag = f"{_A}lvl{level + 1}pPr"
        node = shape
        for _ in range(3):
            if node is None:
                break
            lst = node._element.find(".//" + _A + "lstStyle")
            if lst is not None:
                flag = _bullet_flag(lst.find(tag))
                if flag is not None:
                    return flag
            try:
                node = node._base_placeholder
            except Exception:
                node = None
        try:
            master = self.slide.slide_layout.slide_master._element
        except Exception:
            return None
        body_style = master.find(f"{_P}txStyles/{_P}bodyStyle")
        return _bullet_flag(body_style.find(tag)) if body_style is not None else None

    def table(self, shape) -> None:
        rows: list[list[str]] = []
        for row in shape.table.rows:
            cells: list[str] = []
            for cell in row.cells:
                if getattr(cell, "is_spanned", False):
                    cells.append("")
                    continue
                cells.append(" ".join(_paragraph_text(p._p) for p in cell.text_frame.paragraphs))
            rows.append(cells)
        self.add(table_block(rows))

    def chart(self, shape) -> None:
        text = _alt_text(shape)
        if not text:
            try:
                chart = shape.chart
                if chart.has_title and chart.chart_title.has_text_frame:
                    title = clean_display_text(chart.chart_title.text_frame.text)
                    text = f"Chart: {title}" if title else ""
            except Exception:
                text = ""
        self.add(text_block(BlockKind.FIGURE, text))

    def smartart(self, shape) -> bool:
        """Read SmartArt diagram text as list items. Returns True if the shape was SmartArt."""
        rel_ids = shape._element.find(".//" + _DGM + "relIds")
        if rel_ids is None:
            return False
        try:
            part = self.slide.part.related_part(rel_ids.get(_R + "dm"))
            root = etree.fromstring(part.blob, _XML_PARSER)
        except Exception:
            return True
        for point in root.iter(_DGM + "pt"):
            if point.get("type", "node") not in ("node", "asst"):
                continue
            body = point.find(_DGM + "t")
            if body is None:
                continue
            text = " ".join(_paragraph_text(p) for p in body.iter(_A + "p"))
            self.add(text_block(BlockKind.LIST_ITEM, text, meta={"list_level": 1}))
        return True

    # ------------------------------------------------------------ slide

    def read(self) -> str:
        """Fill ``self.blocks``; returns the slide title ("" if none)."""
        title_shape = None
        try:
            title_shape = self.slide.shapes.title
        except Exception:
            pass
        title = ""
        if title_shape is not None and title_shape.has_text_frame:
            title = clean_display_text(" ".join(_paragraph_text(p._p) for p in title_shape.text_frame.paragraphs))
        if title:
            self.add(text_block(BlockKind.HEADING, title, level=2))
        self.shapes(self.slide.shapes, title_shape)
        if self.slide.has_notes_slide:
            frame = self.slide.notes_slide.notes_text_frame
            if frame is not None:
                for paragraph in frame.paragraphs:
                    self.add(text_block(BlockKind.PARAGRAPH, _paragraph_text(paragraph._p), meta={"speaker_notes": True}))
        return title


def load_pptx(path: Path, *, source_name: str, progress: ProgressFn | None = None) -> Document:
    """Load a PowerPoint ``.pptx`` presentation: one section per slide."""
    import pptx

    path = Path(path)
    if progress:
        progress(0.0, "Opening presentation")
    try:
        presentation = pptx.Presentation(str(path))
    except Exception as exc:
        raise IngestError(
            "This PowerPoint file couldn't be opened. It may be damaged, password-protected, or an older .ppt "
            "file (open it in PowerPoint and use Save As > PowerPoint Presentation (.pptx))."
        ) from exc

    slides = list(presentation.slides)
    sections = []
    first_title = ""
    empty: list[int] = []
    for number, slide in enumerate(slides, start=1):
        if progress:
            progress(0.05 + 0.9 * (number - 1) / max(1, len(slides)), f"Reading slide {number} of {len(slides)}")
        reader = _SlideReader(presentation, slide, number)
        slide_title = reader.read()
        first_title = first_title or slide_title
        if not reader.blocks:
            empty.append(number)
        heading = f"Slide {number}: {slide_title}" if slide_title else f"Slide {number}"
        section = make_section(heading, 1, reader.blocks)
        section.page_start = section.page_end = number
        sections.append(section)

    core_title = core_author = None
    try:
        core_title = clean_display_text(presentation.core_properties.title or "")
        core_author = clean_display_text(presentation.core_properties.author or "") or None
    except Exception:
        pass
    if core_title and core_title.lower() in _DEFAULT_TITLES:
        core_title = None
    title = core_title or first_title or title_from_filename(source_name)

    warnings: list[str] = []
    if not slides:
        warnings.append("This presentation has no slides.")
    elif len(empty) == len(slides):
        warnings.append("No readable text was found in this presentation. Its slides may be pictures only.")
    elif empty:
        label = "Slide" if len(empty) == 1 else "Slides"
        warnings.append(f"{label} {_page_list(empty)} had no text to read (pictures only, or empty).")
    if progress:
        progress(1.0, "Done")
    return Document(
        id=doc_id_for_file(path),
        title=title,
        author=core_author,
        source_name=source_name,
        source_type="pptx",
        page_count=len(slides),
        sections=finalize_sections(sections, title),
        warnings=warnings,
    )
