"""``destiny self-test``: check that everything Athena needs works on this computer.

The packaged app runs it with ``Athena.exe --self-test`` (``--self-test-quick``
for the fast version) and CI runs it on every platform. It prints a readable
report and saves it to ``<data>/logs/self-test.json`` (machine-readable) and
``self-test.txt``. It never sends anything over the internet.

The exit code is 0 when the essentials work: every core module imports, a PDF
and a Word file can be read and planned for speech, the reader's server answers,
and the voices bundled with the app speak. Optional parts that are missing (a
voice engine, OCR, the native window) are warnings.
"""

from __future__ import annotations

import importlib
import importlib.util
import json
import math
import os
import platform
import re
import socket
import sys
import tempfile
import threading
import time
import urllib.request
from array import array
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, TextIO

from . import __version__, config

APP_NAME = "Athena"
PASS, WARN, FAIL, SKIP = "pass", "warn", "fail", "skip"
_MARKS = {PASS: "[ OK ]", WARN: "[WARN]", FAIL: "[FAIL]", SKIP: "[SKIP]"}

TEST_SENTENCE = "Athena is reading this sentence aloud to check that the voice works."
WARM_SENTENCES = (
    "Photosynthesis turns light, water and carbon dioxide into sugar and oxygen.",
    "In 1859, Charles Darwin published On the Origin of Species, which changed biology forever.",
)
COLD_TIMEOUT_S = 240.0
WARM_TIMEOUT_S = 90.0
MIN_SPEECH_S = 0.3
MIN_RMS_DBFS = -50.0

CORE_MODULES = (
    "destiny_bot.models",
    "destiny_bot.config",
    "destiny_bot.ingest",
    "destiny_bot.ingest.common",
    "destiny_bot.ingest.pdf",
    "destiny_bot.ingest.epub",
    "destiny_bot.ingest.html",
    "destiny_bot.ingest.docx",
    "destiny_bot.ingest.pptx",
    "destiny_bot.ingest.text",
    "destiny_bot.ingest.image",
    "destiny_bot.ingest.ocr",
    "destiny_bot.speech",
    "destiny_bot.tts",
    "destiny_bot.tts.catalog",
    "destiny_bot.tts.cache",
    "destiny_bot.tts.piper_engine",
    "destiny_bot.tts.system_engine",
    "destiny_bot.tts.edge_engine",
    "destiny_bot.tts.browser_engine",
    "destiny_bot.audio",
    "destiny_bot.export",
    "destiny_bot.library",
    "destiny_bot.jobs",
    "destiny_bot.ai",
    "destiny_bot.server",
    "destiny_bot.cli",
    "destiny_bot.desktop",
    "pydantic",
    "fastapi",
    "starlette",
    "uvicorn",
    "uvicorn.logging",
    "uvicorn.loops.auto",
    "uvicorn.protocols.http.auto",
    "python_multipart",
    "pymupdf",
    "ebooklib",
    "bs4",
    "lxml.etree",
    "docx",
    "pptx",
    "lameenc",
)
# module -> what it is needed for (missing ones are warnings)
OPTIONAL_MODULES: dict[str, str] = {
    "destiny_bot.tts.kokoro_engine": "Kokoro voices",
    "piper": "natural offline voices (Piper)",
    "onnxruntime": "natural offline voices",
    "numpy": "natural offline voices",
    "kokoro_onnx": "Kokoro voices",
    "edge_tts": "Microsoft online voices",
    "anthropic": "AI study helpers",
    "pytesseract": "reading scanned pages (OCR)",
    "PIL": "reading photos and scanned pages",
    "webview": "the app window",
}
WINDOWS_MODULES: dict[str, str] = {
    "pyttsx3": "Windows voices (SAPI)",
    "comtypes": "Windows voices (SAPI)",
    "clr": "the app window (WebView2)",
}


@dataclass
class Check:
    section: str
    name: str
    status: str
    detail: str = ""
    seconds: float | None = None
    data: dict[str, Any] = field(default_factory=dict)


class Report:
    """Collects checks, printing each one as it completes."""

    def __init__(self, out: TextIO | None) -> None:
        self.out = out
        self.checks: list[Check] = []
        self.lines: list[str] = []
        self.extra: dict[str, Any] = {}

    def say(self, line: str = "") -> None:
        self.lines.append(line)
        if self.out is not None:
            try:
                self.out.write(line + "\n")
                self.out.flush()
            except (OSError, ValueError, UnicodeError):
                pass

    def section(self, title: str) -> None:
        self.say("")
        self.say(title)

    def add(
        self, section: str, name: str, status: str, detail: str = "", *, seconds: float | None = None, **data: Any
    ) -> Check:
        check = Check(section, name, status, detail, round(seconds, 3) if seconds is not None else None, data)
        self.checks.append(check)
        timing = f" ({seconds:.2f} s)" if seconds is not None and seconds >= 0.01 else ""
        text = f"  {_MARKS[status]} {name}{timing}"
        if detail:
            text += f": {detail}"
        self.say(_ascii(text))
        return check

    def count(self, status: str) -> int:
        return sum(1 for c in self.checks if c.status == status)


_ASCII_LOOKALIKES = str.maketrans({"→": "->", "–": "-", "—": "-", "‘": "'", "’": "'", "“": '"', "”": '"', "…": "..."})


def _ascii(text: str) -> str:
    """Console-safe text (Windows consoles may not handle every character)."""
    return text.translate(_ASCII_LOOKALIKES).encode("ascii", "replace").decode("ascii")


def _short_error(exc: BaseException) -> str:
    text = str(exc).strip().splitlines()
    return f"{type(exc).__name__}: {text[0]}" if text else type(exc).__name__


def _run_with_timeout(fn: Callable[[], Any], timeout: float) -> tuple[Any, float]:
    """``(result, seconds)``; raises the function's exception or ``TimeoutError``.

    Runs in a daemon thread so a hung voice engine can't hang the whole self-test.
    """
    box: dict[str, Any] = {}

    def target() -> None:
        try:
            box["result"] = fn()
        except BaseException as exc:  # handed back to the caller below
            box["error"] = exc

    thread = threading.Thread(target=target, name="self-test", daemon=True)
    start = time.perf_counter()
    thread.start()
    thread.join(timeout)
    elapsed = time.perf_counter() - start
    if thread.is_alive():
        raise TimeoutError(f"no answer after {timeout:.0f} seconds")
    if "error" in box:
        raise box["error"]
    return box.get("result"), elapsed


# -- environment


def _memory_gb() -> float | None:
    try:
        if os.name == "nt":
            import ctypes

            class MemoryStatus(ctypes.Structure):
                _fields_ = [
                    ("dwLength", ctypes.c_ulong),
                    ("dwMemoryLoad", ctypes.c_ulong),
                    ("ullTotalPhys", ctypes.c_ulonglong),
                    ("ullAvailPhys", ctypes.c_ulonglong),
                    ("ullTotalPageFile", ctypes.c_ulonglong),
                    ("ullAvailPageFile", ctypes.c_ulonglong),
                    ("ullTotalVirtual", ctypes.c_ulonglong),
                    ("ullAvailVirtual", ctypes.c_ulonglong),
                    ("ullAvailExtendedVirtual", ctypes.c_ulonglong),
                ]

            status = MemoryStatus()
            status.dwLength = ctypes.sizeof(MemoryStatus)
            if ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(status)):
                return round(status.ullTotalPhys / 1024**3, 1)
            return None
        return round(os.sysconf("SC_PHYS_PAGES") * os.sysconf("SC_PAGE_SIZE") / 1024**3, 1)
    except (AttributeError, OSError, ValueError):
        return None


def _environment() -> dict[str, Any]:
    from . import desktop

    info: dict[str, Any] = {
        "app": APP_NAME,
        "version": __version__,
        "frozen": desktop.is_frozen(),
        "executable": sys.executable,
        "python": {
            "version": platform.python_version(),
            "implementation": platform.python_implementation(),
            "bits": 64 if sys.maxsize > 2**32 else 32,
            "build": " ".join(platform.python_build()),
        },
        "os": {
            "system": platform.system(),
            "release": platform.release(),
            "version": platform.version(),
            "platform": platform.platform(),
            "machine": platform.machine(),
        },
        "cpu_count": os.cpu_count(),
        "memory_gb": _memory_gb(),
        "data_dir": str(config.data_dir()),
        "extra_voice_dirs": os.environ.get("DESTINY_EXTRA_VOICE_DIRS", ""),
    }
    if sys.platform == "win32":
        release, version, csd, _ptype = platform.win32_ver()
        info["os"]["windows"] = {"release": release, "build": version, "service_pack": csd}
        try:
            info["os"]["windows"]["edition"] = platform.win32_edition()
        except AttributeError:
            pass
    elif sys.platform == "darwin":
        info["os"]["macos"] = platform.mac_ver()[0]
    else:
        try:
            info["os"]["distro"] = platform.freedesktop_os_release().get("PRETTY_NAME")
        except (OSError, AttributeError):
            pass
    bundled = desktop.bundled_voices_dir()
    if bundled is not None:
        info["bundled_voices"] = sorted(
            f"{p.relative_to(bundled).as_posix()} ({p.stat().st_size / 1e6:.1f} MB)"
            for p in bundled.rglob("*")
            if p.is_file()
        )
    return info


# -- checks


def _check_imports(report: Report) -> None:
    report.section("Program parts")
    failed: list[str] = []
    start = time.perf_counter()
    for name in CORE_MODULES:
        try:
            importlib.import_module(name)
        except Exception as exc:
            failed.append(name)
            report.add("imports", name, FAIL, _short_error(exc))
    if not failed:
        report.add("imports", f"{len(CORE_MODULES)} core modules", PASS, seconds=time.perf_counter() - start)
    optional = dict(OPTIONAL_MODULES)
    if sys.platform == "win32":
        optional.update(WINDOWS_MODULES)
    for name, purpose in optional.items():
        try:
            if name.startswith("destiny_bot.") and importlib.util.find_spec(name) is None:
                report.add("imports", name, SKIP, f"not part of this version ({purpose})")
                continue
            begin = time.perf_counter()
            importlib.import_module(name)
            report.add("imports", name, PASS, seconds=time.perf_counter() - begin)
        except Exception as exc:
            report.add("imports", name, WARN, f"{_short_error(exc)} (needed for {purpose})")


def _check_window(report: Report) -> None:
    from . import desktop

    report.section("App window")
    data: dict[str, Any] = {}
    if sys.platform == "win32":
        version = desktop._webview2_version()
        data["webview2"] = version
        report.add(
            "window",
            "Microsoft Edge WebView2 runtime",
            PASS if version else WARN,
            f"version {version}" if version else "not installed; Athena will open in an Edge app window instead",
        )
    browser = desktop.find_app_browser()
    data["app_browser"] = str(browser) if browser else None
    report.add(
        "window",
        "Edge/Chrome for the fallback app window",
        PASS if browser else WARN,
        str(browser) if browser else "not found; the default browser would be used as a last resort",
    )
    report.extra["window"] = data


def _wav_stats(clip: Any) -> dict[str, Any]:
    """Duration, loudness and format of a synthesized clip (raises ValueError if it isn't usable audio)."""
    from .audio import read_wav, to_int16

    if clip.mime != "audio/wav":
        return {"mime": clip.mime, "bytes": len(clip.data)}
    pcm, rate, channels, width = read_wav(clip.data)
    if width != 2:
        pcm = to_int16(pcm, width)
    samples = array("h")
    samples.frombytes(pcm[: len(pcm) - len(pcm) % 2])
    if sys.byteorder == "big":
        samples.byteswap()
    frames = len(samples) // max(channels, 1)
    duration = frames / rate if rate else 0.0
    peak = max((abs(s) for s in samples), default=0)
    rms = math.sqrt(sum(s * s for s in samples) / len(samples)) if samples else 0.0
    rms_dbfs = 20 * math.log10(rms / 32768) if rms > 0 else -120.0
    return {
        "mime": clip.mime,
        "bytes": len(clip.data),
        "sample_rate": rate,
        "channels": channels,
        "duration_s": round(duration, 3),
        "peak": peak,
        "rms_dbfs": round(rms_dbfs, 1),
    }


def _valid_speech(stats: dict[str, Any]) -> str:
    """'' if the clip sounds like speech, else what is wrong with it."""
    if stats.get("mime") != "audio/wav":
        return "" if stats.get("bytes", 0) > 1000 else "the audio file is empty"
    if stats["duration_s"] < MIN_SPEECH_S:
        return f"only {stats['duration_s']:.2f} s of audio"
    if stats["peak"] < 500 or stats["rms_dbfs"] < MIN_RMS_DBFS:
        return f"the audio is silent (level {stats['rms_dbfs']} dBFS)"
    return ""


def _bundled_engines() -> set[str]:
    """Engines whose voices ship with the app: they must work."""
    from . import desktop

    folder = desktop.bundled_voices_dir()
    names: set[str] = set()
    if folder is None:
        return names
    if any((folder / "piper").glob("*.onnx")):
        names.add("piper")
    if any((folder / "kokoro").glob("*.onnx")):
        names.add("kokoro")
    return names


def _check_engines(report: Report, quick: bool) -> None:
    report.section("Voices")
    try:
        from . import tts

        engines = list(tts.all_engines())
    except Exception as exc:
        report.add("engines", "voice engines", FAIL, _short_error(exc))
        return
    bundled = _bundled_engines()
    infos: list[dict[str, Any]] = []
    synth_results: list[dict[str, Any]] = []
    for engine in engines:
        info: dict[str, Any] = {"name": engine.name, "label": engine.label, "online": engine.online}
        try:
            (ok, reason), seconds = _run_with_timeout(engine.is_available, 120)
            info.update(available=bool(ok), reason=reason, check_seconds=round(seconds, 3))
            if ok:
                voices, list_seconds = _run_with_timeout(engine.list_voices, 120)
                info["voices"] = [v.id for v in voices][:50]
                info["voice_count"] = len(voices)
                info["list_seconds"] = round(list_seconds, 3)
                info["default_voice"] = engine.default_voice()
        except Exception as exc:
            info.update(available=False, reason=_short_error(exc))
        infos.append(info)
        must_work = engine.name in bundled
        if info.get("available"):
            detail = f"{info.get('voice_count', 0)} voice(s), default {info.get('default_voice') or 'none'}"
            report.add("engines", f"{engine.label}", PASS, detail, seconds=info.get("check_seconds"))
        else:
            status = FAIL if must_work else WARN
            report.add("engines", f"{engine.label}", status, info.get("reason") or "not available")

    report.extra["engines"] = infos
    working = 0
    for engine, info in zip(engines, infos):
        if not info.get("available") or engine.online or engine.name == "browser":
            continue
        result = _benchmark(engine, info.get("default_voice"), quick)
        synth_results.append(result)
        must_work = engine.name in bundled
        if result.get("error"):
            report.add("synthesis", f"{engine.name} speaks", FAIL if must_work else WARN, result["error"])
            continue
        working += 1
        detail = (
            f"voice {result['voice']}: {result['audio']['duration_s']:.1f} s of audio, "
            f"first {result['cold_s']:.2f} s (RTF {result['cold_rtf']:.2f})"
        )
        if "warm_s" in result:
            detail += f", then {result['warm_s']:.2f} s (RTF {result['warm_rtf']:.2f})"
        report.add("synthesis", f"{engine.name} speaks", PASS, detail, seconds=result["cold_s"])
    for name in sorted(bundled - {e.name for e in engines}):
        report.add("synthesis", f"{name} speaks", FAIL, "voices for it are bundled, but the engine is missing")
    if working == 0:
        report.add(
            "synthesis",
            "an offline voice",
            WARN,
            "no offline voice could speak; the reader will fall back to the browser's built-in voice",
        )
    for engine, info in zip(engines, infos):
        if engine.online:
            report.add("synthesis", f"{engine.name} speaks", SKIP, "online voice: not tested, to keep this check offline")
    report.extra["synthesis"] = synth_results


def _benchmark(engine: Any, voice: str | None, quick: bool) -> dict[str, Any]:
    """Speak a test sentence: cold (includes loading the voice) and, unless quick, warm runs."""
    result: dict[str, Any] = {"engine": engine.name, "voice": voice}
    try:
        clip, cold = _run_with_timeout(lambda: engine.synthesize(TEST_SENTENCE, voice), COLD_TIMEOUT_S)
        stats = _wav_stats(clip)
        problem = _valid_speech(stats)
        result.update(cold_s=round(cold, 3), audio=stats)
        if problem:
            result["error"] = problem
            return result
        duration = stats.get("duration_s") or 0.0
        result["cold_rtf"] = round(cold / duration, 3) if duration else None
        if quick:
            return result
        runs: list[tuple[float, float]] = []
        for sentence in WARM_SENTENCES:
            warm_clip, seconds = _run_with_timeout(
                lambda text=sentence: engine.synthesize(text, voice), WARM_TIMEOUT_S
            )
            warm_stats = _wav_stats(warm_clip)
            if _valid_speech(warm_stats):
                result["error"] = f"a second sentence failed: {_valid_speech(warm_stats)}"
                return result
            runs.append((seconds, warm_stats.get("duration_s") or 0.0))
        best = min(runs, key=lambda r: r[0] / r[1] if r[1] else math.inf)
        result["warm_s"] = round(best[0], 3)
        result["warm_audio_s"] = round(best[1], 3)
        result["warm_rtf"] = round(best[0] / best[1], 3) if best[1] else None
        result["warm_runs"] = [{"seconds": round(s, 3), "audio_s": round(a, 3)} for s, a in runs]
    except Exception as exc:
        result["error"] = _short_error(exc)
    return result


def _registry_voices() -> dict[str, list[dict[str, str]]]:
    """Voices registered with Windows: SAPI 5 (what Athena's system voice uses) and OneCore."""
    import winreg

    places = {
        "sapi5": (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Speech\Voices\Tokens"),
        "sapi5_32bit": (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\WOW6432Node\Microsoft\Speech\Voices\Tokens"),
        "sapi5_user": (winreg.HKEY_CURRENT_USER, r"SOFTWARE\Microsoft\Speech\Voices\Tokens"),
        "onecore": (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Speech_OneCore\Voices\Tokens"),
    }
    found: dict[str, list[dict[str, str]]] = {}
    for label, (root, path) in places.items():
        voices: list[dict[str, str]] = []
        try:
            with winreg.OpenKey(root, path) as tokens:
                index = 0
                while True:
                    try:
                        token = winreg.EnumKey(tokens, index)
                    except OSError:
                        break
                    index += 1
                    entry = {"token": token}
                    try:
                        with winreg.OpenKey(tokens, token) as key:
                            entry["name"] = str(winreg.QueryValueEx(key, "")[0])
                        with winreg.OpenKey(tokens, token + r"\Attributes") as attrs:
                            for attr in ("Language", "Gender", "Age", "Vendor"):
                                try:
                                    entry[attr.lower()] = str(winreg.QueryValueEx(attrs, attr)[0])
                                except OSError:
                                    pass
                    except OSError:
                        pass
                    voices.append(entry)
        except OSError:
            pass
        found[label] = voices
    return found


def _check_windows_voices(report: Report) -> None:
    report.section("Windows voices")
    try:
        registry = _registry_voices()
    except Exception as exc:
        registry = {}
        report.add("sapi", "voice registry", WARN, _short_error(exc))
    sapi_names = [v.get("name") or v["token"] for v in registry.get("sapi5", []) + registry.get("sapi5_user", [])]
    onecore_names = [v.get("name") or v["token"] for v in registry.get("onecore", [])]
    report.add(
        "sapi",
        "SAPI 5 voices",
        PASS if sapi_names else WARN,
        ", ".join(sapi_names) if sapi_names else "none registered",
    )
    report.add(
        "sapi",
        "OneCore voices (Settings > Speech)",
        PASS if onecore_names else SKIP,
        ", ".join(onecore_names) if onecore_names else "none",
    )
    try:
        from .tts import get_engine

        voices, seconds = _run_with_timeout(get_engine("system").list_voices, 120)
        listed = [f"{v.id} ({v.language or '?'})" for v in voices]
        report.add(
            "sapi",
            "voices the system engine sees",
            PASS if listed else WARN,
            ", ".join(listed) if listed else "none",
            seconds=seconds,
        )
    except Exception as exc:
        listed = []
        report.add("sapi", "voices the system engine sees", WARN, _short_error(exc))
    report.extra["windows_voices"] = {"registry": registry, "system_engine": listed}


def _make_pdf(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    chapters = [
        ("Chapter 1: Cells", "Cells are the basic units of life. Every living thing is made of one or more cells. "
         "The cell membrane controls what enters and leaves the cell."),
        ("Chapter 2: Energy", "Mitochondria release energy from food. Plants also capture energy from sunlight, "
         "for example in the leaves."),
    ]  # fmt: skip
    for title, body in chapters:
        page = doc.new_page()
        page.insert_text((72, 90), title, fontsize=20)
        page.insert_textbox(pymupdf.Rect(72, 120, 520, 400), body, fontsize=11)
    doc.save(str(path))
    doc.close()


def _make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Chapter 1: Cells", level=1)
    document.add_paragraph("Cells are the basic units of life. Every living thing is made of one or more cells.")
    document.add_paragraph("The nucleus holds the cell's DNA, e.g. its genes.")
    document.add_heading("Chapter 2: Energy", level=1)
    document.add_paragraph("Mitochondria release energy from food.")
    document.save(str(path))


def _check_ingest(report: Report) -> None:
    report.section("Reading files")
    from .ingest import load_document
    from .speech import plan_section

    with tempfile.TemporaryDirectory(prefix="athena-selftest-") as tmp:
        for kind, maker in (("PDF", _make_pdf), ("Word", _make_docx)):
            path = Path(tmp) / ("sample.pdf" if kind == "PDF" else "sample.docx")
            try:
                start = time.perf_counter()
                maker(path)
                doc = load_document(path, ocr=False)
                section = next((s for s in doc.sections if s.blocks), None)
                if section is None:
                    raise ValueError("no readable text was found")
                text = " ".join(b.text for s in doc.sections for b in s.blocks).lower()
                if "cells" not in text:
                    raise ValueError("the text came out wrong")
                plan = plan_section(section)
                if not plan or any(not u.speech.strip() for u in plan):
                    raise ValueError("the reading plan was empty")
                report.add(
                    "ingest",
                    f"{kind} file",
                    PASS,
                    f"{len(doc.sections)} section(s); '{section.title}' -> {len(plan)} sentence(s) to read",
                    seconds=time.perf_counter() - start,
                )
            except Exception as exc:
                report.add("ingest", f"{kind} file", FAIL, _short_error(exc))


def _free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def _http_get(url: str, headers: dict[str, str] | None = None) -> tuple[int, bytes]:
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    with opener.open(urllib.request.Request(url, headers=headers or {}), timeout=30) as response:
        return response.status, response.read()


def _check_server(report: Report) -> None:
    report.section("Reader")
    from .server import start_server_in_thread

    stop: Callable[[], None] | None = None
    data: dict[str, Any] = {}
    try:
        (url, stop), seconds = _run_with_timeout(lambda: start_server_in_thread(port=_free_port()), 60)
        report.add("server", "start the reader's server", PASS, url, seconds=seconds)
        begin = time.perf_counter()
        status, body = _http_get(url)
        page = body.decode("utf-8", "replace")
        match = re.search(r'<meta\s+name="destiny-token"\s+content="([^"]+)"', page)
        if status != 200 or not match:
            raise ValueError(f"the reader page didn't load (HTTP {status})")
        report.add("server", "GET /", PASS, f"{len(body) // 1024} KB page", seconds=time.perf_counter() - begin)
        begin = time.perf_counter()
        status, body = _http_get(url.rstrip("/") + "/api/status", {"X-Destiny-Token": match.group(1)})
        payload = json.loads(body)
        engines = {e.get("name"): bool(e.get("available")) for e in payload.get("engines", [])}
        data = {"url": url, "resolved_engine": payload.get("resolved_engine"), "engines": engines,
                "version": payload.get("version"), "ocr": payload.get("ocr")}  # fmt: skip
        report.add(
            "server",
            "GET /api/status",
            PASS,
            f"voice engine in use: {payload.get('resolved_engine')}",
            seconds=time.perf_counter() - begin,
        )
    except Exception as exc:
        report.add("server", "the reader's server", FAIL, _short_error(exc))
    finally:
        if stop is not None:
            try:
                _, seconds = _run_with_timeout(stop, 30)
                report.add("server", "stop the server", PASS, seconds=seconds)
            except Exception as exc:
                report.add("server", "stop the server", WARN, _short_error(exc))
    report.extra["server"] = data


def _check_ocr(report: Report, quick: bool) -> None:
    report.section("Scanned pages (OCR)")
    try:
        from .ingest.ocr import ocr_available, ocr_image_bytes

        ok, reason = ocr_available()
    except Exception as exc:
        report.add("ocr", "OCR", WARN, _short_error(exc))
        return
    if not ok:
        report.add("ocr", "OCR", WARN, reason or "not available")
        return
    if quick:
        report.add("ocr", "OCR", PASS, "available")
        return
    try:
        import io

        from PIL import Image, ImageDraw, ImageFont

        image = Image.new("L", (900, 160), 255)
        try:
            font: Any = ImageFont.load_default(size=56)
        except TypeError:  # Pillow < 10.1
            font = ImageFont.load_default()
        ImageDraw.Draw(image).text((30, 40), "Athena reads textbooks", fill=0, font=font)
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        text, seconds = _run_with_timeout(lambda: ocr_image_bytes(buf.getvalue()), 120)
        if "textbooks" in text.lower():
            report.add("ocr", "OCR reads a test image", PASS, seconds=seconds)
        else:
            report.add("ocr", "OCR reads a test image", WARN, f"read {text.strip()[:60]!r}", seconds=seconds)
    except Exception as exc:
        report.add("ocr", "OCR reads a test image", WARN, _short_error(exc))


# -- running


def _write_atomic(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=f".{path.name}-", suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as fh:
            fh.write(text)
        os.replace(tmp, path)
    except BaseException:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def report_paths() -> tuple[Path, Path]:
    """Where the JSON and text reports are saved."""
    logs = config.data_dir() / "logs"
    return logs / "self-test.json", logs / "self-test.txt"


def run_self_test(*, quick: bool = False, out: TextIO | None = None) -> int:
    """Run every check, print the report and save it; 0 when the essentials work."""
    out = out if out is not None else sys.stdout
    report = Report(out)
    started = time.perf_counter()
    report.say(_ascii(f"{APP_NAME} self-test, version {__version__}{' (quick)' if quick else ''}"))
    try:
        from . import desktop

        desktop.use_bundled_voices()
    except Exception as exc:
        report.add("environment", "bundled voices", WARN, _short_error(exc))
    try:
        environment = _environment()
    except Exception as exc:
        environment = {"error": _short_error(exc)}
    os_info = environment.get("os", {})
    report.say(
        _ascii(
            f"Python {environment.get('python', {}).get('version', '?')} on {os_info.get('platform', '?')}; "
            f"data folder {environment.get('data_dir', '?')}"
        )
    )

    steps: list[tuple[str, Callable[[], None]]] = [
        ("imports", lambda: _check_imports(report)),
        ("window", lambda: _check_window(report)),
        ("engines", lambda: _check_engines(report, quick)),
    ]
    if sys.platform == "win32":
        steps.append(("sapi", lambda: _check_windows_voices(report)))
    steps += [
        ("ingest", lambda: _check_ingest(report)),
        ("server", lambda: _check_server(report)),
        ("ocr", lambda: _check_ocr(report, quick)),
    ]
    for name, step in steps:
        try:
            step()
        except Exception as exc:  # a broken check is itself a failure, never a crash
            report.add(name, f"{name} check", FAIL, _short_error(exc))

    total = time.perf_counter() - started
    failed, warned, passed = report.count(FAIL), report.count(WARN), report.count(PASS)
    ok = failed == 0
    report.say("")
    if ok:
        verdict = "Everything Athena needs is working" + (f" ({warned} optional part(s) missing)." if warned else ".")
    else:
        verdict = f"{failed} problem(s) found. Athena may not work properly."
    report.say(f"{verdict} {passed} passed, {warned} warning(s), {failed} failed, {total:.1f} s.")

    json_path, text_path = report_paths()
    payload = {
        "app": APP_NAME,
        "version": __version__,
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "quick": quick,
        "ok": ok,
        "summary": {"pass": passed, "warn": warned, "fail": failed, "skip": report.count(SKIP)},
        "seconds_total": round(total, 2),
        "environment": environment,
        "checks": [asdict(c) for c in report.checks],
        **report.extra,
    }
    try:
        _write_atomic(json_path, json.dumps(payload, indent=2, default=str))
        report.say(_ascii(f"Report saved to {json_path}"))
        _write_atomic(text_path, "\n".join(report.lines) + "\n")
    except OSError as exc:
        report.say(_ascii(f"Couldn't save the report: {exc}"))
    return 0 if ok else 1


__all__ = ["Check", "Report", "report_paths", "run_self_test"]
