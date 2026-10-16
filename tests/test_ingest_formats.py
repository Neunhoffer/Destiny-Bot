"""Tests for the non-PDF loaders (HTML, EPUB, Word, PowerPoint, text, images) and OCR.

Every fixture is generated in code: EPUBs with ebooklib, Word files with
python-docx, slides with python-pptx and pictures with Pillow.
"""

from __future__ import annotations

import copy
import importlib
import io
import re
import sys
import zipfile
from pathlib import Path

import pytest

from athena.ingest import IngestError, load_document
from athena.ingest import document_from_text as dispatch_document_from_text
from athena.ingest import ocr as ocr_module
from athena.ingest.html import html_to_blocks, load_html, parse_html
from athena.ingest.text import (
    decode_text_bytes,
    document_from_text,
    looks_like_markdown,
    markdown_blocks,
    plain_text_blocks,
)
from athena.models import Block, BlockKind, Document

K = BlockKind
_MARKUP = re.compile(r"</?[A-Za-z][A-Za-z0-9:-]*(?:\s[^<>]*)?/?>")
_MD_MARKUP = re.compile(r"\*\*|__|\]\(|^#{1,6}\s|^>\s|```")


# --------------------------------------------------------------------------- helpers


def all_blocks(doc: Document) -> list[Block]:
    return [block for section in doc.sections for block in section.blocks]


def pairs(blocks: list[Block]) -> list[tuple[BlockKind, str]]:
    return [(block.kind, block.text) for block in blocks]


def texts(doc: Document, kind: BlockKind | None = None) -> list[str]:
    return [b.text for b in all_blocks(doc) if kind is None or b.kind == kind]


def assert_sane(doc: Document) -> None:
    """Structure every loader must guarantee."""
    assert doc.sections, "a document always has at least one section"
    assert [s.id for s in doc.sections] == [f"s{i}" for i in range(len(doc.sections))]
    assert doc.title.strip()
    assert re.fullmatch(r"[0-9a-f]{16}", doc.id)
    for section in doc.sections:
        assert section.title.strip() and section.title == section.title.strip()
        assert section.level >= 1
        for block in section.blocks:
            assert block.text.strip() or block.rows, f"empty block in {section.title!r}"
            assert block.text == block.text.strip()
            assert not _MARKUP.search(block.text), f"stray markup: {block.text!r}"
            if block.kind != K.CODE:
                assert "\n" not in block.text and "  " not in block.text, repr(block.text)
            if block.kind == K.HEADING:
                assert block.level is not None and 1 <= block.level <= 6
            if block.kind == K.TABLE:
                assert block.rows and all(isinstance(cell, str) for row in block.rows for cell in row)
        if section.blocks and section.blocks[0].kind == K.HEADING and section.level == section.blocks[0].level:
            assert section.title == section.blocks[0].text


def png_bytes(size=(80, 60), color="navy") -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, "PNG")
    return buffer.getvalue()


def _ocr_ready() -> bool:
    ocr_module.reset_ocr_cache()
    ok = ocr_module.ocr_available()[0]
    ocr_module.reset_ocr_cache()
    return ok


requires_tesseract = pytest.mark.skipif(not _ocr_ready(), reason="Tesseract OCR is not installed")


@pytest.fixture
def fresh_ocr_cache():
    ocr_module.reset_ocr_cache()
    yield
    ocr_module.reset_ocr_cache()


# --------------------------------------------------------------------------- HTML

PAGE = """<!DOCTYPE html>
<html><head><title>Cell Biology &amp; More</title><meta name="author" content="Ada Lovelace">
<style>p { color: red }</style><script>var x = "<p>not text</p>";</script></head>
<body>
<header role="banner"><a href="/">Campus Site</a><nav><a href="/a">Home</a></nav></header>
<nav><ul><li>Menu item</li></ul></nav>
<main>
<article>
<header><h1 id="top">Chapter 1: <em>Cells</em></h1></header>
<p>The <span>cell</span> is the <b>basic</b> unit<sup><a href="#fn1" id="r1">1</a></sup> of life.
Inline math <math alttext="x^2"><msup><mi>x</mi><mn>2</mn></msup></math> stays inline. E = mc<sup>2</sup>.</p>
<div>Text directly in a div <span>with a span</span>
  <p>A child paragraph.</p>
  trailing div text
</div>
<p hidden>Hidden paragraph</p>
<p aria-hidden="true">Aria hidden paragraph</p>
<p style="display: none">Display none paragraph</p>
<template><p>Template text</p></template>
<noscript><p>Enable JavaScript</p></noscript>
<svg><title>Svg title</title><text x="1" y="1">Svg label</text></svg>
<h2>Lists</h2>
<ol start="3"><li>Third item<ul><li>Nested bullet</li></ul></li><li><p>Fourth as a paragraph</p></li></ol>
<blockquote><p>To be or not to be.</p><p>Second quote paragraph.</p></blockquote>
<figure><img src="c.png" alt="Diagram of a cell"><figcaption>Figure 1.1 A typical cell.</figcaption></figure>
<p><img src="decor.png" alt=""></p>
<table><caption>Table 1.1 Organelles</caption><thead><tr><th>Organelle</th><th>Function</th></tr></thead>
<tbody><tr><td>Nucleus</td><td>Stores <i>DNA</i></td></tr><tr><td>Ribosome</td><td>Makes proteins</td></tr></tbody></table>
<pre><code>def double(x):
    return x * 2</code></pre>
<p><code>pip install athena</code></p>
<div class="eq"><math display="block"><semantics><mrow><mi>a</mi><mo>=</mo><mfrac><mi>b</mi><mi>c</mi></mfrac></mrow>
<annotation encoding="application/x-tex">a = \\frac{b}{c}</annotation></semantics></math></div>
<math display="block"><mrow><mi>y</mi><mo>=</mo><msqrt><mi>x</mi></msqrt><mo>+</mo><mn>1</mn></mrow></math>
<p>Page break<span epub:type="pagebreak" id="p12" title="12">12</span> marker removed.</p>
<footer><p>Posted in Biology</p></footer>
</article>
</main>
<section role="doc-endnotes"><hr><ol><li id="fn1"><p>First footnote text. <a href="#r1" role="doc-backlink">&#x21A9;&#xFE0E;</a></p></li></ol></section>
<footer role="contentinfo">Copyright 2026 Campus Inc.</footer>
</body></html>"""


def test_html_to_blocks_maps_structure():
    blocks, title = html_to_blocks(PAGE)
    assert title == "Cell Biology & More"
    assert pairs(blocks) == [
        (K.HEADING, "Chapter 1: Cells"),
        (K.PARAGRAPH, "The cell is the basic unit of life. Inline math x^2 stays inline. E = mc²."),
        (K.PARAGRAPH, "Text directly in a div with a span"),
        (K.PARAGRAPH, "A child paragraph."),
        (K.PARAGRAPH, "trailing div text"),
        (K.HEADING, "Lists"),
        (K.LIST_ITEM, "3. Third item"),
        (K.LIST_ITEM, "Nested bullet"),
        (K.LIST_ITEM, "4. Fourth as a paragraph"),
        (K.QUOTE, "To be or not to be."),
        (K.QUOTE, "Second quote paragraph."),
        (K.FIGURE, "Diagram of a cell"),
        (K.CAPTION, "Figure 1.1 A typical cell."),
        (K.CAPTION, "Table 1.1 Organelles"),
        (K.TABLE, "Organelle | Function; Nucleus | Stores DNA; Ribosome | Makes proteins"),
        (K.CODE, "def double(x):\n    return x * 2"),
        (K.CODE, "pip install athena"),
        (K.EQUATION, "a = \\frac{b}{c}"),
        (K.EQUATION, "y = sqrt(x) + 1"),
        (K.PARAGRAPH, "Page break marker removed."),
        (K.PARAGRAPH, "Posted in Biology"),
        (K.FOOTNOTE, "1. First footnote text."),
    ]
    assert [b.level for b in blocks if b.kind == K.HEADING] == [1, 2]
    assert next(b for b in blocks if b.kind == K.TABLE).rows == [
        ["Organelle", "Function"],
        ["Nucleus", "Stores DNA"],
        ["Ribosome", "Makes proteins"],
    ]
    nested = next(b for b in blocks if b.text == "Nested bullet")
    assert nested.meta["list_level"] == 2
    assert all(b.page is None for b in blocks)


def test_html_page_numbers_and_parse_details():
    content = parse_html(PAGE, page=7)
    assert content.author == "Ada Lovelace"
    assert all(b.page == 7 for b in content.blocks)
    assert content.anchors["top"] == 0


def test_html_div_text_and_inline_spans_do_not_split():
    html = (
        "<body><div><span>One</span> <span>sentence</span> <a href='x'>with</a> <b>many</b> spans.</div>"
        "<div><div>Nested <i>div</i> text.</div></div><section>Loose section text.<p>Para.</p></section>"
        "<p>She said <q>cells divide</q>.</p></body>"
    )
    blocks, _ = html_to_blocks(html)
    assert pairs(blocks) == [
        (K.PARAGRAPH, "One sentence with many spans."),
        (K.PARAGRAPH, "Nested div text."),
        (K.PARAGRAPH, "Loose section text."),
        (K.PARAGRAPH, "Para."),
        (K.PARAGRAPH, "She said “cells divide”."),
    ]


def test_html_footnotes_epub_and_pandoc_styles():
    html = """<html xmlns:epub="http://www.idpf.org/2007/ops"><body>
    <p>Cells divide<a epub:type="noteref" href="#n1">1</a> and grow<a role="doc-noteref" href="#n2">2</a>.
    Pandoc ref<a href="#fn3" class="footnote-ref" id="fnref3"><sup>3</sup></a> here.</p>
    <aside epub:type="footnote" id="n1"><p>Mitosis makes two cells.</p></aside>
    <div role="doc-endnote" id="n2"><p>Growth needs energy.</p></div>
    <p class="footnote">A classic footnote paragraph.</p>
    <section class="footnotes footnotes-end-of-document" role="doc-endnotes"><hr />
    <ol><li id="fn3"><p>Pandoc note.<a href="#fnref3" class="footnote-back" role="doc-backlink">&#x21A9;&#xFE0E;</a></p></li></ol></section>
    </body></html>"""
    blocks, _ = html_to_blocks(html)
    assert pairs(blocks) == [
        (K.PARAGRAPH, "Cells divide and grow. Pandoc ref here."),
        (K.FOOTNOTE, "Mitosis makes two cells."),
        (K.FOOTNOTE, "Growth needs energy."),
        (K.FOOTNOTE, "A classic footnote paragraph."),
        (K.FOOTNOTE, "1. Pandoc note."),
    ]


def test_html_xhtml_self_closing_tags_do_not_swallow_text():
    xhtml = """<?xml version="1.0" encoding="utf-8"?>
    <html xmlns="http://www.w3.org/1999/xhtml"><head><title>X</title><script src="a.js"/></head>
    <body><p><a id="page5"/>First paragraph.</p><div/><p>Second <span class="x"/>paragraph.</p></body></html>"""
    blocks, title = html_to_blocks(xhtml.encode("utf-8"))
    assert title == "X"
    assert pairs(blocks) == [(K.PARAGRAPH, "First paragraph."), (K.PARAGRAPH, "Second paragraph.")]


def test_html_layout_tables_are_read_as_content():
    html = """<body><table><tr><td><h2>Welcome</h2><p>Layout cell text.</p></td></tr></table>
    <table role="presentation"><tr><td>Left</td><td>Right</td></tr></table></body>"""
    blocks, _ = html_to_blocks(html)
    assert pairs(blocks) == [
        (K.HEADING, "Welcome"),
        (K.PARAGRAPH, "Layout cell text."),
        (K.PARAGRAPH, "Left"),
        (K.PARAGRAPH, "Right"),
    ]


def test_html_mathml_without_alttext_is_linearised():
    html = """<body><p><math display="block"><mrow><msup><mi>x</mi><mn>2</mn></msup><mo>+</mo>
    <mfrac><mrow><mi>a</mi><mo>+</mo><mi>b</mi></mrow><mn>2</mn></mfrac><mo>=</mo>
    <msub><mi>y</mi><mn>0</mn></msub></mrow></math></p></body>"""
    blocks, _ = html_to_blocks(html)
    assert pairs(blocks) == [(K.EQUATION, "x^2 + (a + b)/2 = y_0")]
    # Placeholder alt text is ignored; wide spaces separate side-by-side expressions.
    html = """<body><math alttext="Alternative text not available"><mtable><mtr><mtd><mi>x</mi></mtd>
    <mtd><mo>=</mo><mn>1</mn><mspace width="2em"/></mtd><mtd><mi>y</mi></mtd><mtd><mo>=</mo><mn>2</mn></mtd>
    </mtr></mtable></math><div class="related-links"><a href="next.html">Next topic</a></div></body>"""
    blocks, _ = html_to_blocks(html)
    assert pairs(blocks) == [(K.EQUATION, "x = 1; y = 2")]


def test_html_charset_detection():
    declared = b'<html><head><meta charset="windows-1252"></head><body><p>\x93Quoted\x94 caf\xe9</p></body></html>'
    blocks, _ = html_to_blocks(declared)
    assert blocks[0].text == "“Quoted” café"
    latin1_label = b'<html><head><meta http-equiv="Content-Type" content="text/html; charset=iso-8859-1"></head>'
    blocks, _ = html_to_blocks(latin1_label + b"<body><p>\x93Smart\x94</p></body></html>")
    assert blocks[0].text == "“Smart”"  # browsers treat latin-1 labels as Windows-1252
    blocks, _ = html_to_blocks("<p>Grüße — naïve</p>".encode())
    assert blocks[0].text == "Grüße — naïve"
    blocks, _ = html_to_blocks("<p>UTF-16 text</p>".encode("utf-16"))
    assert blocks[0].text == "UTF-16 text"


def test_html_main_content_drops_site_chrome_but_keeps_notes():
    html = """<body><div class="topbar">Log in | Sign up | Search the whole campus site for news</div>
    <main><h1>Photosynthesis</h1><p>Plants convert light into chemical energy.<sup><a href="#n1">1</a></sup></p></main>
    <aside epub:type="footnote" id="n1"><p>Mostly in chloroplasts.</p></aside></body>"""
    blocks, _ = html_to_blocks(html)
    assert pairs(blocks) == [
        (K.HEADING, "Photosynthesis"),
        (K.PARAGRAPH, "Plants convert light into chemical energy."),
        (K.FOOTNOTE, "Mostly in chloroplasts."),
    ]


def test_html_old_style_breaks_and_external_note_links():
    html = """<body>First line<br>same paragraph<br><br>Second paragraph<sup><a href="notes.xhtml#n4">4</a></sup>.
    <p>Linked<a href="#f2"><sup>2</sup></a> marker and <a href="other.html#sec">a real link</a>.</p>
    <p><a href="big.png"><img alt="Plant cell"/> <img alt="Animal cell"/></a></p></body>"""
    blocks, _ = html_to_blocks(html)
    assert pairs(blocks) == [
        (K.PARAGRAPH, "First line same paragraph"),
        (K.PARAGRAPH, "Second paragraph."),
        (K.PARAGRAPH, "Linked marker and a real link."),
        (K.FIGURE, "Plant cell"),
        (K.FIGURE, "Animal cell"),
    ]


def test_html_absurdly_deep_nesting_falls_back_to_flat_reading():
    html = "<body><h1>Title</h1>" + "<div>" * 3000 + "<p>Deep text.</p><p>More <b>bold</b> text.</p>" + "</div>" * 3000
    blocks, _ = html_to_blocks(html + "</body>")
    assert pairs(blocks) == [(K.HEADING, "Title"), (K.PARAGRAPH, "Deep text."), (K.PARAGRAPH, "More bold text.")]


def test_load_html_title_sections_and_dispatch(tmp_path):
    page = tmp_path / "lecture_notes.html"
    page.write_text(PAGE, encoding="utf-8")
    doc = load_document(page)
    assert_sane(doc)
    assert doc.source_type == "html" and doc.source_name == "lecture_notes.html"
    assert doc.title == "Cell Biology & More"
    assert doc.author == "Ada Lovelace"
    assert [(s.title, s.level) for s in doc.sections] == [("Chapter 1: Cells", 1), ("Lists", 2)]
    assert doc.sections[0].blocks[0].kind == K.HEADING
    assert doc == load_html(page, source_name="lecture_notes.html")

    untitled = tmp_path / "untitled.htm"
    untitled.write_bytes(b"<body><h1>First Heading</h1><p>Body.</p></body>")
    assert load_document(untitled).title == "First Heading"
    bare = tmp_path / "week_3_reading.xhtml"
    bare.write_bytes(b"<body><p>Just text.</p></body>")
    doc = load_document(bare)
    assert doc.title == "week 3 reading"
    assert [s.title for s in doc.sections] == ["week 3 reading"]

    empty = tmp_path / "empty.html"
    empty.write_bytes(b"<html><body><script>x()</script></body></html>")
    doc = load_document(empty)
    assert_sane(doc)
    assert doc.warnings


# --------------------------------------------------------------------------- EPUB


def _xhtml(title: str, body: str) -> str:
    return f"<html><head><title>{title}</title></head><body>{body}</body></html>"


def build_epub(path: Path, *, with_nav: bool = True, with_toc: bool = True) -> Path:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("urn:uuid:athena-test")
    book.set_title("Introductory Biology")
    book.set_language("en")
    book.add_author("Jane Smith")
    book.add_author("Raj Patel")
    book.set_cover("images/cover.png", png_bytes((60, 90)))

    ch1 = epub.EpubHtml(title="Chapter 1", file_name="text/ch1.xhtml", lang="en")
    ch1.content = _xhtml(
        "Chapter 1",
        '<section epub:type="chapter"><h1>Chapter 1: Cells</h1>'
        '<p>All living things are made of cells.<a epub:type="noteref" href="#fn1">1</a> '
        "Cells were named by Robert Hooke.</p>"
        '<aside epub:type="footnote" id="fn1"><p>Hooke looked at cork in 1665.</p></aside>'
        '<h2 id="membranes">1.2 Membranes</h2><p>Membranes control what enters the cell.</p></section>',
    )
    ch2 = epub.EpubHtml(title="Chapter 2", file_name="text/ch2.xhtml", lang="en")
    ch2.content = _xhtml(
        "Chapter 2",
        "<h1>Chapter 2: Energy</h1>"
        '<p>Energy is <math xmlns="http://www.w3.org/1998/Math/MathML" alttext="E = m c^2">'
        "<mi>E</mi><mo>=</mo><mi>m</mi><msup><mi>c</mi><mn>2</mn></msup></math> in physics.</p>"
        '<math xmlns="http://www.w3.org/1998/Math/MathML" display="block" alttext="\\Delta G = \\Delta H - T \\Delta S">'
        "<mi>G</mi></math>"
        "<table><tr><th>Molecule</th><th>Energy</th></tr><tr><td>ATP</td><td>High</td></tr></table>"
        '<figure><img src="../images/cover.png" alt="Diagram of ATP"/><figcaption>Figure 2.1 ATP</figcaption></figure>',
    )
    ch2b = epub.EpubHtml(title="Chapter 2 continued", file_name="text/ch2b.xhtml", lang="en")
    ch2b.content = _xhtml("Chapter 2", "<p>Continuation of chapter two without its own heading.</p>")
    ch3 = epub.EpubHtml(title="Chapter 3", file_name="text/ch3.xhtml", lang="en")
    ch3.content = _xhtml(
        "Chapter 3",
        "<h1>Chapter 3: Genetics</h1><ol><li>Genes are made of DNA.</li><li>DNA is copied.</li></ol>"
        "<blockquote><p>Like begets like.</p></blockquote>",
    )
    for item in (ch1, ch2, ch2b, ch3):
        book.add_item(item)
    if with_toc:
        book.toc = [
            (
                epub.Section("Chapter 1: Cells", href="text/ch1.xhtml"),
                [epub.Link("text/ch1.xhtml#membranes", "1.2 Membranes", "m")],
            ),
            epub.Link("text/ch2.xhtml", "Chapter 2: Energy", "c2"),
            epub.Link("text/ch3.xhtml", "Chapter 3: Genetics", "c3"),
        ]
    book.add_item(epub.EpubNcx())
    spine = ["cover", ch1, ch2, ch2b, ch3]
    if with_nav:
        book.add_item(epub.EpubNav())
        spine.insert(1, "nav")
    book.spine = spine
    epub.write_epub(str(path), book)
    return path


def test_epub_uses_toc_sections_and_html_mapping(tmp_path):
    path = build_epub(tmp_path / "bio.epub")
    doc = load_document(path)
    assert_sane(doc)
    assert doc.source_type == "epub"
    assert doc.title == "Introductory Biology"
    assert doc.author == "Jane Smith, Raj Patel"
    assert [(s.title, s.level) for s in doc.sections] == [
        ("Chapter 1: Cells", 1),
        ("1.2 Membranes", 2),
        ("Chapter 2: Energy", 1),
        ("Chapter 3: Genetics", 1),
    ]
    ch1, membranes, ch2, ch3 = doc.sections
    assert pairs(ch1.blocks) == [
        (K.HEADING, "Chapter 1: Cells"),
        (K.PARAGRAPH, "All living things are made of cells. Cells were named by Robert Hooke."),
        (K.FOOTNOTE, "Hooke looked at cork in 1665."),
    ]
    assert pairs(membranes.blocks)[0] == (K.HEADING, "1.2 Membranes")
    assert pairs(ch2.blocks) == [
        (K.HEADING, "Chapter 2: Energy"),
        (K.PARAGRAPH, "Energy is E = m c^2 in physics."),
        (K.EQUATION, "\\Delta G = \\Delta H - T \\Delta S"),
        (K.TABLE, "Molecule | Energy; ATP | High"),
        (K.FIGURE, "Diagram of ATP"),
        (K.CAPTION, "Figure 2.1 ATP"),
        (K.PARAGRAPH, "Continuation of chapter two without its own heading."),
    ]
    assert pairs(ch3.blocks) == [
        (K.HEADING, "Chapter 3: Genetics"),
        (K.LIST_ITEM, "1. Genes are made of DNA."),
        (K.LIST_ITEM, "2. DNA is copied."),
        (K.QUOTE, "Like begets like."),
    ]
    all_text = " ".join(texts(doc))
    assert "Cover" not in all_text  # image-only cover page skipped
    assert "Table of Contents" not in all_text and doc.warnings == []


def test_epub2_ncx_table_of_contents(tmp_path):
    doc = load_document(build_epub(tmp_path / "ncx.epub", with_nav=False))
    assert_sane(doc)
    assert [s.title for s in doc.sections] == [
        "Chapter 1: Cells",
        "1.2 Membranes",
        "Chapter 2: Energy",
        "Chapter 3: Genetics",
    ]


def test_epub_without_toc_splits_by_headings(tmp_path):
    doc = load_document(build_epub(tmp_path / "notoc.epub", with_toc=False))
    assert_sane(doc)
    assert [(s.title, s.level) for s in doc.sections] == [
        ("Chapter 1: Cells", 1),
        ("1.2 Membranes", 2),
        ("Chapter 2: Energy", 1),
        ("Chapter 3: Genetics", 1),
    ]
    assert doc.sections[2].blocks[-1].text == "Continuation of chapter two without its own heading."


def _write_epub_with_encryption(path: Path, algorithm: str, source: Path | None = None) -> Path:
    encryption = (
        '<?xml version="1.0"?><encryption xmlns="urn:oasis:names:tc:opendocument:xmlns:container" '
        'xmlns:enc="http://www.w3.org/2001/04/xmlenc#"><enc:EncryptedData>'
        f'<enc:EncryptionMethod Algorithm="{algorithm}"/>'
        '<enc:CipherData><enc:CipherReference URI="EPUB/text/ch1.xhtml"/></enc:CipherData>'
        "</enc:EncryptedData></encryption>"
    )
    with zipfile.ZipFile(path, "w") as out:
        if source is None:
            out.writestr("mimetype", "application/epub+zip")
            out.writestr(
                "META-INF/container.xml",
                '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                "</rootfiles></container>",
            )
            out.writestr("OEBPS/content.opf", "scrambled")
        else:
            with zipfile.ZipFile(source) as original:
                for info in original.infolist():
                    out.writestr(info, original.read(info))
        out.writestr("META-INF/encryption.xml", encryption)
    return path


def test_epub_drm_is_refused(tmp_path):
    path = _write_epub_with_encryption(tmp_path / "locked.epub", "http://www.w3.org/2001/04/xmlenc#aes128-cbc")
    with pytest.raises(IngestError, match="DRM-protected"):
        load_document(path)


def test_epub_font_obfuscation_is_not_drm(tmp_path):
    source = build_epub(tmp_path / "plain.epub")
    path = _write_epub_with_encryption(tmp_path / "fonts.epub", "http://www.idpf.org/2008/embedding", source)
    doc = load_document(path)
    assert doc.sections[0].title == "Chapter 1: Cells"


def test_epub_broken_file_gives_friendly_error(tmp_path):
    bad = tmp_path / "broken.epub"
    bad.write_bytes(b"this is not a zip file")
    with pytest.raises(IngestError, match="couldn't be opened"):
        load_document(bad)


# --------------------------------------------------------------------------- DOCX

_OMML_DISPLAY = (
    '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath>'
    "<m:r><m:t>E</m:t></m:r><m:r><m:t>=</m:t></m:r><m:r><m:t>m</m:t></m:r>"
    "<m:sSup><m:e><m:r><m:t>c</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>"
    "</m:oMath></m:oMathPara>"
)
_OMML_INLINE = (
    '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
    "<m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b+1</m:t></m:r></m:den></m:f></m:oMath>"
)


def build_docx(path: Path) -> Path:
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Inches

    d = docx.Document()
    d.core_properties.title = "Core Title"
    d.core_properties.author = "Prof. Ada Lovelace"
    d.add_paragraph("Cell Biology Handout", style="Title")
    d.add_paragraph("Week 3 reading", style="Subtitle")
    d.add_heading("Introduction", level=1)

    p = d.add_paragraph("Cells are the basic unit of life. Read ")
    rid = d.part.relate_to("https://example.edu/cells", RT.HYPERLINK, is_external=True)
    p._p.append(
        parse_xml(
            f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}"><w:r><w:t xml:space="preserve">the course page</w:t>'
            "</w:r></w:hyperlink>"
        )
    )
    p.add_run(" before class.")
    d.add_paragraph("")
    d.add_paragraph("   ")

    p = d.add_paragraph("Einstein wrote E = mc")
    p.add_run("2").font.superscript = True
    p.add_run(" HIDDEN NOTE").font.hidden = True
    p.add_run(".")

    p = d.add_paragraph("Mitochondria ")
    p._p.append(
        parse_xml(f'<w:del {nsdecls("w")} w:id="1" w:author="x"><w:r><w:delText>destroy</w:delText></w:r></w:del>')
    )
    p._p.append(parse_xml(f'<w:ins {nsdecls("w")} w:id="2" w:author="x"><w:r><w:t>make</w:t></w:r></w:ins>'))
    p.add_run(" energy.")

    d.add_paragraph("Diffusion needs no energy", style="List Bullet")
    d.add_paragraph("Facilitated diffusion uses proteins", style="List Bullet 2")
    d.add_paragraph("Collect the sample", style="List Number")

    table = d.add_table(rows=3, cols=2)
    for r, row in enumerate([("Organelle", "Function"), ("Nucleus", "Stores DNA"), ("Ribosome", "Makes proteins")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    d.add_paragraph("Nothing in biology makes sense except in the light of evolution.", style="Quote")
    d.add_heading("Equations", level=2)
    d.add_paragraph()._p.append(parse_xml(_OMML_DISPLAY))
    p = d.add_paragraph("The ratio ")
    p._p.append(parse_xml(_OMML_INLINE))
    p.add_run(" stays constant.")

    footnotes = (
        f"<w:footnotes {nsdecls('w')}>"
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r>'
        '<w:r><w:t xml:space="preserve"> Hooke first saw cells in cork.</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    )
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        footnotes.encode("utf-8"),
        d.part.package,
    )
    d.part.relate_to(part, RT.FOOTNOTES)
    p = d.add_paragraph("Robert Hooke named them cells.")
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:footnoteReference w:id="1"/></w:r>'))

    d.add_picture(io.BytesIO(png_bytes()), width=Inches(1))
    doc_pr = d.paragraphs[-1]._p.xpath(".//wp:docPr")[0]
    doc_pr.set("descr", "Microscope image of onion cells")

    wps = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    textbox = (
        f'<w:r {nsdecls("w", "wp", "a")} xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
        f' xmlns:wps="{wps}" xmlns:v="urn:schemas-microsoft-com:vml"><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:anchor>'
        '<wp:docPr id="9" name="Text Box 9"/><a:graphic><a:graphicData uri="x"><wps:wsp><wps:txbx><w:txbxContent>'
        "<w:p><w:r><w:t>Sidebar: cells were discovered in 1665.</w:t></w:r></w:p>"
        "</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        "<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r>"
        "<w:t>Sidebar: cells were discovered in 1665.</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>"
        "</mc:Fallback></mc:AlternateContent></w:r>"
    )
    p = d.add_paragraph("See the box.")
    p._p.append(parse_xml(textbox))

    d.add_heading("Summary", level=1)
    d.add_paragraph("Cells are amazing.")
    d.save(str(path))
    return path


def test_docx_structure(tmp_path):
    doc = load_document(build_docx(tmp_path / "handout.docx"))
    assert_sane(doc)
    assert doc.source_type == "docx"
    assert doc.title == "Cell Biology Handout"
    assert doc.author == "Prof. Ada Lovelace"
    assert [(s.title, s.level) for s in doc.sections] == [
        ("Cell Biology Handout", 1),
        ("Introduction", 1),
        ("Equations", 2),
        ("Summary", 1),
    ]
    blocks = all_blocks(doc)
    assert pairs(blocks) == [
        (K.HEADING, "Cell Biology Handout"),
        (K.PARAGRAPH, "Week 3 reading"),
        (K.HEADING, "Introduction"),
        (K.PARAGRAPH, "Cells are the basic unit of life. Read the course page before class."),
        (K.PARAGRAPH, "Einstein wrote E = mc²."),
        (K.PARAGRAPH, "Mitochondria make energy."),
        (K.LIST_ITEM, "Diffusion needs no energy"),
        (K.LIST_ITEM, "Facilitated diffusion uses proteins"),
        (K.LIST_ITEM, "Collect the sample"),
        (K.TABLE, "Organelle | Function; Nucleus | Stores DNA; Ribosome | Makes proteins"),
        (K.QUOTE, "Nothing in biology makes sense except in the light of evolution."),
        (K.HEADING, "Equations"),
        (K.EQUATION, "E=mc^2"),
        (K.PARAGRAPH, "The ratio a/(b+1) stays constant."),
        (K.PARAGRAPH, "Robert Hooke named them cells."),
        (K.FOOTNOTE, "Hooke first saw cells in cork."),
        (K.FIGURE, "Microscope image of onion cells"),
        (K.PARAGRAPH, "See the box."),
        (K.PARAGRAPH, "Sidebar: cells were discovered in 1665."),
        (K.HEADING, "Summary"),
        (K.PARAGRAPH, "Cells are amazing."),
    ]
    headings = [b for b in blocks if b.kind == K.HEADING]
    assert [h.level for h in headings] == [1, 1, 2, 1]
    lists = [b for b in blocks if b.kind == K.LIST_ITEM]
    assert [b.meta["list_level"] for b in lists] == [1, 2, 1]
    table = next(b for b in blocks if b.kind == K.TABLE)
    assert table.rows == [["Organelle", "Function"], ["Nucleus", "Stores DNA"], ["Ribosome", "Makes proteins"]]


def test_docx_custom_styles_numbering_and_code(tmp_path):
    import docx
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    d = docx.Document()
    chapter = d.styles.add_style("Chapter Heading", WD_STYLE_TYPE.PARAGRAPH)
    chapter.base_style = d.styles["Heading 1"]
    d.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    code = d.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    code.base_style = d.styles["Normal"]
    d.add_paragraph("Contents entry ........ 3", style="TOC 1")
    d.add_paragraph("Chapter One", style="Chapter Heading")
    numbered = d.add_paragraph("Numbered without a list style")
    numbered._p.get_or_add_pPr().append(
        parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="1"/><w:numId w:val="1"/></w:numPr>')
    )
    d.add_paragraph("def area(r):", style="Source Code")
    d.add_paragraph("    return 3.14 * r * r", style="Source Code")
    d.add_paragraph("Back to prose.")
    d.save(str(tmp_path / "styles.docx"))
    doc = load_document(tmp_path / "styles.docx")
    assert_sane(doc)
    blocks = all_blocks(doc)
    assert pairs(blocks) == [
        (K.HEADING, "Chapter One"),
        (K.LIST_ITEM, "Numbered without a list style"),
        (K.CODE, "def area(r):\n    return 3.14 * r * r"),
        (K.PARAGRAPH, "Back to prose."),
    ]
    assert blocks[0].level == 1 and blocks[1].meta["list_level"] == 2


def test_docx_title_falls_back_to_core_properties_then_filename(tmp_path):
    import docx

    d = docx.Document()
    d.core_properties.title = "Lab Safety Rules"
    d.add_paragraph("Wear goggles.")
    d.save(str(tmp_path / "rules.docx"))
    assert load_document(tmp_path / "rules.docx").title == "Lab Safety Rules"

    d = docx.Document()
    d.add_paragraph("Wear goggles.")
    d.save(str(tmp_path / "lab_rules.docx"))
    doc = load_document(tmp_path / "lab_rules.docx")
    assert doc.title == "lab rules"
    assert_sane(doc)


def test_docx_broken_file(tmp_path):
    bad = tmp_path / "old.docx"
    bad.write_bytes(b"\xd0\xcf\x11\xe0 not really a docx")
    with pytest.raises(IngestError, match="Word file couldn't be opened"):
        load_document(bad)


# --------------------------------------------------------------------------- PPTX


def build_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.util import Inches

    prs = Presentation()
    prs.core_properties.title = "Cells Lecture"
    prs.core_properties.author = "Dr. Rivera"

    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Intro to Cells"
    s1.placeholders[1].text = "Biology 101"

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Key Ideas"
    body = s2.placeholders[1].text_frame
    body.text = "Cells are the unit of life"
    for text, level in (("Prokaryotes lack a nucleus", 1), ("Eukaryotes have organelles", 0)):
        para = body.add_paragraph()
        para.text = text
        para.level = level
    s2.notes_slide.notes_text_frame.text = "Remember to mention mitochondria.\nAsk a question."

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Data"
    table = s3.shapes.add_table(3, 2, Inches(1), Inches(1.8), Inches(6), Inches(1.5)).table
    for r, row in enumerate([("Cell", "Size"), ("Bacteria", "1 µm"), ("Neuron", "100 µm")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    chart_data = CategoryChartData()
    chart_data.categories = ["Plant", "Animal"]
    chart_data.add_series("Count", (3, 5))
    chart = s3.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(4), Inches(3), chart_data
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Cell counts"

    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    s4.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(5))._element.nvPicPr.cNvPr.set(
        "descr", "Microscope photo of cheek cells"
    )
    group = s4.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(5), Inches(2), Inches(3), Inches(1)).text_frame.text = "Right column"
    group.shapes.add_textbox(Inches(0.5), Inches(2.1), Inches(3), Inches(1)).text_frame.text = "Left column"
    s4.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(1)).text_frame.text = "Top note"
    for placeholder in s4.slide_layout.placeholders:
        if placeholder.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            s4.shapes._spTree.append(copy.deepcopy(placeholder._element))

    prs.slides.add_slide(prs.slide_layouts[6])  # an empty slide
    prs.save(str(path))
    return path


def test_pptx_slides_become_sections(tmp_path):
    doc = load_document(build_pptx(tmp_path / "lecture.pptx"))
    assert_sane(doc)
    assert doc.source_type == "pptx"
    assert doc.title == "Cells Lecture"
    assert doc.author == "Dr. Rivera"
    assert doc.page_count == 5
    assert [(s.title, s.level, s.page_start, s.page_end) for s in doc.sections] == [
        ("Slide 1: Intro to Cells", 1, 1, 1),
        ("Slide 2: Key Ideas", 1, 2, 2),
        ("Slide 3: Data", 1, 3, 3),
        ("Slide 4", 1, 4, 4),
    ]
    s1, s2, s3, s4 = doc.sections
    assert pairs(s1.blocks) == [(K.HEADING, "Intro to Cells"), (K.PARAGRAPH, "Biology 101")]
    assert s1.blocks[0].level == 2
    assert pairs(s2.blocks) == [
        (K.HEADING, "Key Ideas"),
        (K.LIST_ITEM, "Cells are the unit of life"),
        (K.LIST_ITEM, "Prokaryotes lack a nucleus"),
        (K.LIST_ITEM, "Eukaryotes have organelles"),
        (K.PARAGRAPH, "Remember to mention mitochondria."),
        (K.PARAGRAPH, "Ask a question."),
    ]
    assert [b.meta.get("list_level") for b in s2.blocks[1:4]] == [1, 2, 1]
    assert [bool(b.meta.get("speaker_notes")) for b in s2.blocks] == [False] * 4 + [True] * 2
    assert pairs(s3.blocks) == [
        (K.HEADING, "Data"),
        (K.TABLE, "Cell | Size; Bacteria | 1 µm; Neuron | 100 µm"),
        (K.FIGURE, "Chart: Cell counts"),
    ]
    assert pairs(s4.blocks) == [
        (K.PARAGRAPH, "Top note"),
        (K.PARAGRAPH, "Left column"),
        (K.PARAGRAPH, "Right column"),
        (K.FIGURE, "Microscope photo of cheek cells"),
    ]
    assert all(b.page == number for number, sec in enumerate(doc.sections, 1) for b in sec.blocks)
    assert doc.warnings == ["Slide 5 had no text to read (pictures only, or empty)."]


def test_pptx_explicit_bullets_override_defaults(tmp_path):
    from lxml import etree
    from pptx import Presentation
    from pptx.util import Inches

    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Bullets"
    body = slide.placeholders[1].text_frame
    body.text = "Plain statement in a body placeholder"
    body.paragraphs[0]._p.get_or_add_pPr().append(etree.SubElement(etree.Element("x"), f"{{{a}}}buNone"))
    box = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(1)).text_frame
    box.text = "Text box with a bullet"
    char = etree.Element(f"{{{a}}}buChar")
    char.set("char", "•")
    box.paragraphs[0]._p.get_or_add_pPr().append(char)
    prs.save(str(tmp_path / "bullets.pptx"))
    doc = load_document(tmp_path / "bullets.pptx")
    assert pairs(doc.sections[0].blocks) == [
        (K.HEADING, "Bullets"),
        (K.PARAGRAPH, "Plain statement in a body placeholder"),
        (K.LIST_ITEM, "Text box with a bullet"),
    ]


def test_pptx_title_from_first_slide_and_broken_file(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Photosynthesis"
    prs.save(str(tmp_path / "deck.pptx"))
    doc = load_document(tmp_path / "deck.pptx")
    assert doc.title == "Photosynthesis"
    assert [s.title for s in doc.sections] == ["Slide 1: Photosynthesis"]

    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"PK\x03\x04 truncated")
    with pytest.raises(IngestError, match="PowerPoint file couldn't be opened"):
        load_document(bad)


# --------------------------------------------------------------------------- text and Markdown

MARKDOWN = """---
title: Front Matter Title
author: Jane Doe
---
# Biology Notes

Intro with **bold**, *italic*, `code*x*`, a [link](http://x.com "t"), snake_case_name and 2 * 3 * 4.
Second line with ![inline icon](i.png) image.

Setext Heading
--------------

- item one
- item two
  continued here
  - nested item
* [x] done task
1. first
2. second

> A quote line
> continues here
>
> Second quote para

```python
def f():
    return 1
```

| Name | Age |
|------|----:|
| Alice | 30 |
| Bob \\| Jr | 25 |

![A diagram of a cell](cell.png)

<!-- a comment
spanning lines -->

Text with a footnote[^1] and \\*escaped\\* stars &amp; an entity.

[^1]: The footnote text.

***

$$
E = mc^2
$$
"""


def test_markdown_blocks():
    assert looks_like_markdown(MARKDOWN)
    doc = document_from_text(MARKDOWN, "")
    assert_sane(doc)
    assert doc.title == "Front Matter Title"
    assert doc.author == "Jane Doe"
    assert [(s.title, s.level) for s in doc.sections] == [("Biology Notes", 1), ("Setext Heading", 2)]
    assert pairs(all_blocks(doc)) == [
        (K.HEADING, "Biology Notes"),
        (
            K.PARAGRAPH,
            "Intro with bold, italic, code*x*, a link, snake_case_name and 2 * 3 * 4. "
            "Second line with inline icon image.",
        ),
        (K.HEADING, "Setext Heading"),
        (K.LIST_ITEM, "item one"),
        (K.LIST_ITEM, "item two continued here"),
        (K.LIST_ITEM, "nested item"),
        (K.LIST_ITEM, "done task"),
        (K.LIST_ITEM, "1. first"),
        (K.LIST_ITEM, "2. second"),
        (K.QUOTE, "A quote line continues here"),
        (K.QUOTE, "Second quote para"),
        (K.CODE, "def f():\n    return 1"),
        (K.TABLE, "Name | Age; Alice | 30; Bob | Jr | 25"),
        (K.FIGURE, "A diagram of a cell"),
        (K.PARAGRAPH, "Text with a footnote and *escaped* stars & an entity."),
        (K.FOOTNOTE, "The footnote text."),
        (K.EQUATION, "E = mc^2"),
    ]
    table = next(b for b in all_blocks(doc) if b.kind == K.TABLE)
    assert table.rows == [["Name", "Age"], ["Alice", "30"], ["Bob | Jr", "25"]]
    code = next(b for b in all_blocks(doc) if b.kind == K.CODE)
    assert code.meta == {"language": "python"}
    assert [b.meta["list_level"] for b in all_blocks(doc) if b.kind == K.LIST_ITEM] == [1, 1, 2, 1, 1, 1]


def test_markdown_indented_code_and_html_lines():
    blocks = markdown_blocks(
        '# T\n\nPara one.\n\n    indented code\n    more\n\n<div align="center">\n\nCentered text\n\n</div>\n'
    )
    assert pairs(blocks) == [
        (K.HEADING, "T"),
        (K.PARAGRAPH, "Para one."),
        (K.CODE, "indented code\nmore"),
        (K.PARAGRAPH, "Centered text"),
    ]


PLAIN = """CHAPTER 1
The Cell

Cells are the basic unit of life. They were first
described by Robert Hooke in 1665, who saw them in
cork. Infor-
mation flows from DNA.

1.2 Membrane Transport
Membranes control what enters and leaves the cell.

Key points:
- Diffusion needs no energy
• Active transport uses ATP and
  moves against the gradient
1. First step
2. Second step

INTRODUCTION TO GENETICS

PART II

Chapter 3: Cells

Unit 4 Review

Chapter 3 discusses how cells divide

Notes
=====

The end.
"""


def test_plain_text_headings_paragraphs_and_bullets():
    blocks = plain_text_blocks(PLAIN)
    assert pairs(blocks) == [
        (K.HEADING, "CHAPTER 1: The Cell"),
        (
            K.PARAGRAPH,
            "Cells are the basic unit of life. They were first described by Robert Hooke in 1665, "
            "who saw them in cork. Information flows from DNA.",
        ),
        (K.HEADING, "1.2 Membrane Transport"),
        (K.PARAGRAPH, "Membranes control what enters and leaves the cell."),
        (K.PARAGRAPH, "Key points:"),
        (K.LIST_ITEM, "Diffusion needs no energy"),
        (K.LIST_ITEM, "Active transport uses ATP and moves against the gradient"),
        (K.LIST_ITEM, "1. First step"),
        (K.LIST_ITEM, "2. Second step"),
        (K.HEADING, "INTRODUCTION TO GENETICS"),
        (K.HEADING, "PART II"),
        (K.HEADING, "Chapter 3: Cells"),
        (K.HEADING, "Unit 4 Review"),
        (K.PARAGRAPH, "Chapter 3 discusses how cells divide"),
        (K.HEADING, "Notes"),
        (K.PARAGRAPH, "The end."),
    ]
    levels = {b.text: b.level for b in blocks if b.kind == K.HEADING}
    assert levels["PART II"] == 1
    assert levels["Chapter 3: Cells"] == levels["CHAPTER 1: The Cell"] == 2
    assert levels["1.2 Membrane Transport"] == 3


def test_plain_text_unwrapped_lines_are_paragraphs():
    long_a = "This paragraph was typed without hard wrapping, so it is one very long line. " * 4
    long_b = "The next paragraph follows directly on the next line without a blank line between. " * 4
    blocks = plain_text_blocks(f"{long_a}\n{long_b}")
    assert [b.kind for b in blocks] == [K.PARAGRAPH, K.PARAGRAPH]


def test_text_file_encodings(tmp_path):
    smart = "\u201cSmart quotes\u201d \u2014 caf\u00e9 \u2018fine\u2019"
    cases = {
        "cp1252.txt": smart.encode("cp1252"),
        "utf16.txt": smart.encode("utf-16"),
        "utf16be.txt": b"\xfe\xff" + smart.encode("utf-16-be"),
        "utf8sig.txt": smart.encode("utf-8-sig"),
        "utf32.txt": smart.encode("utf-32"),
        "utf8.txt": smart.encode("utf-8"),
        "utf16-nobom.txt": ("Plain ASCII words in a file. " + smart).encode("utf-16-le"),
    }
    for name, data in cases.items():
        path = tmp_path / name
        path.write_bytes(data)
        doc = load_document(path)
        assert_sane(doc)
        assert texts(doc)[-1].endswith(smart), name
        assert doc.source_type == "txt"
    assert decode_text_bytes(b"caf\xe9 \x81") == "caf\xe9 \x81"  # not cp1252 either -> Latin-1


def test_text_file_windows_line_endings_and_title(tmp_path):
    path = tmp_path / "bio_notes.txt"
    path.write_bytes(b"CHAPTER 1\r\n\r\nLine one of a\r\nwrapped paragraph.\r\n\r\nSecond paragraph.\r\n")
    doc = load_document(path)
    assert_sane(doc)
    assert doc.title == "bio notes"
    assert pairs(all_blocks(doc)) == [
        (K.HEADING, "CHAPTER 1"),
        (K.PARAGRAPH, "Line one of a wrapped paragraph."),
        (K.PARAGRAPH, "Second paragraph."),
    ]
    assert [s.title for s in doc.sections] == ["CHAPTER 1"]


def test_markdown_file_title_and_type(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text(
        "# Cell Theory\n\nAll living things are made of cells.\n\n## History\n\nHooke, 1665.\n", encoding="utf-8"
    )
    doc = load_document(path)
    assert_sane(doc)
    assert doc.source_type == "md" and doc.title == "Cell Theory"
    assert [s.title for s in doc.sections] == ["Cell Theory", "History"]
    other = tmp_path / "notes.markdown"
    other.write_text("Just one line of *markdown*.", encoding="utf-8")
    assert texts(load_document(other)) == ["Just one line of markdown."]
    # A .txt file that clearly uses Markdown is parsed as Markdown.
    txt = tmp_path / "readme.txt"
    txt.write_text("# Title\n\nSome **bold** text.\n\n## Part\n\n- a\n- b\n", encoding="utf-8")
    doc = load_document(txt)
    assert doc.source_type == "txt"
    assert texts(doc) == ["Title", "Some bold text.", "Part", "a", "b"]


def test_empty_and_binary_text_files_raise(tmp_path):
    empty = tmp_path / "empty.txt"
    empty.write_bytes(b"  \r\n ")
    with pytest.raises(IngestError, match="empty"):
        load_document(empty)
    binary = tmp_path / "binary.txt"
    binary.write_bytes(bytes(range(1, 32)) * 50)
    with pytest.raises(IngestError, match="doesn't look like a text file"):
        load_document(binary)


def test_document_from_text_paste():
    doc = dispatch_document_from_text("First paragraph.\n\nSecond paragraph.", "My notes")
    assert_sane(doc)
    assert doc.source_type == "paste" and doc.title == "My notes"
    assert texts(doc) == ["First paragraph.", "Second paragraph."]
    same = document_from_text("First paragraph.\n\nSecond paragraph.", "My notes")
    assert same.id == doc.id
    assert document_from_text("First paragraph.\n\nSecond paragraph.", "Other title").id != doc.id
    assert document_from_text("Different text.", "My notes").id != doc.id

    md = document_from_text("# Heading\n\nSome **bold** words and a [link](https://x.y).\n\n## Next\n\nMore.", "Pasted")
    assert texts(md) == ["Heading", "Some bold words and a link.", "Next", "More."]
    assert [s.title for s in md.sections] == ["Heading", "Next"]

    untitled = document_from_text("A short first line\n\nBody text here.", "  ")
    assert untitled.title == "A short first line"
    with pytest.raises(IngestError):
        document_from_text("   \n  ", "Nothing")


# --------------------------------------------------------------------------- OCR and images


def render_text_image(lines: list[str], *, size: int = 40, width: int = 1300):
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.load_default(size=size)
    line_height = int(size * 1.6)
    image = Image.new("RGB", (width, 60 + line_height * len(lines)), "white")
    draw = ImageDraw.Draw(image)
    for i, line in enumerate(lines):
        if line:
            draw.text((40, 30 + i * line_height), line, fill="black", font=font)
    return image


def _image_bytes(image, fmt: str = "PNG") -> bytes:
    buffer = io.BytesIO()
    image.save(buffer, fmt)
    return buffer.getvalue()


PAGE_LINES = [
    "CHAPTER 1",
    "",
    "The cell is the basic unit of life.",
    "All living things are made of cells.",
    "",
    "- Membranes control transport",
]


@requires_tesseract
def test_ocr_image_bytes_reads_rendered_text(fresh_ocr_cache):
    ok, reason = ocr_module.ocr_available()
    assert ok and "Tesseract" in reason
    text = ocr_module.ocr_image_bytes(_image_bytes(render_text_image(PAGE_LINES)))
    assert "CHAPTER 1" in text
    assert "basic unit of life" in text
    # Small screenshots are upscaled, transparent images go on white paper.
    from PIL import Image, ImageDraw, ImageFont

    small = Image.new("RGBA", (420, 60), (0, 0, 0, 0))
    ImageDraw.Draw(small).text(
        (10, 15), "Tiny transparent text", fill=(0, 0, 0, 255), font=ImageFont.load_default(size=18)
    )
    assert "Tiny transparent text" in ocr_module.ocr_image_bytes(_image_bytes(small))


@requires_tesseract
def test_load_image_builds_document(tmp_path, fresh_ocr_cache):
    path = tmp_path / "handout_photo.png"
    path.write_bytes(_image_bytes(render_text_image(PAGE_LINES)))
    doc = load_document(path)
    assert_sane(doc)
    assert doc.source_type == "image" and doc.title == "handout photo"
    assert doc.page_count == 1 and doc.warnings == []
    blocks = all_blocks(doc)
    assert blocks[0].kind == K.HEADING and blocks[0].text == "CHAPTER 1"
    assert blocks[1].kind == K.PARAGRAPH
    assert blocks[1].text.startswith("The cell is the basic unit of life. All living things")
    assert blocks[-1].kind == K.LIST_ITEM and "Membranes control transport" in blocks[-1].text
    assert doc.sections[0].title == "CHAPTER 1"


@requires_tesseract
def test_multipage_tiff_gets_a_section_per_page(tmp_path, fresh_ocr_cache):
    pages = [
        render_text_image(["First page about cells."]).convert("L"),
        render_text_image(["Second page about energy."]).convert("L"),
    ]
    path = tmp_path / "scan.tiff"
    pages[0].save(path, save_all=True, append_images=pages[1:])
    doc = load_document(path)
    assert_sane(doc)
    assert doc.page_count == 2
    assert [s.title for s in doc.sections] == ["Page 1", "Page 2"]
    assert "cells" in doc.sections[0].blocks[0].text and "energy" in doc.sections[1].blocks[0].text


@requires_tesseract
def test_blank_image_returns_document_with_warning(tmp_path, fresh_ocr_cache):
    from PIL import Image

    path = tmp_path / "blank.jpg"
    Image.new("RGB", (800, 600), "white").save(path, "JPEG")
    doc = load_document(path)
    assert_sane(doc)
    assert doc.warnings and "No text was found" in doc.warnings[0]
    assert all_blocks(doc) == []


@requires_tesseract
def test_corrupt_image_raises(tmp_path, fresh_ocr_cache):
    with pytest.raises(IngestError, match="couldn't be opened"):
        ocr_module.ocr_image_bytes(b"definitely not an image")
    bad = tmp_path / "bad.png"
    bad.write_bytes(b"\x89PNG\r\n\x1a\n broken")
    with pytest.raises(IngestError):
        load_document(bad)


def test_ocr_unavailable_when_tesseract_missing(tmp_path, monkeypatch, fresh_ocr_cache):
    pytesseract = pytest.importorskip("pytesseract")
    monkeypatch.setattr(pytesseract.pytesseract, "tesseract_cmd", str(tmp_path / "no-such-tesseract"))
    ok, reason = ocr_module.ocr_available()
    assert not ok and "isn't installed" in reason
    with pytest.raises(IngestError, match="Tesseract"):
        ocr_module.ocr_image_bytes(png_bytes())
    image = tmp_path / "photo.png"
    image.write_bytes(png_bytes())
    with pytest.raises(IngestError, match="needs OCR"):
        load_document(image)


def test_ocr_unavailable_without_pytesseract(monkeypatch, fresh_ocr_cache):
    monkeypatch.setitem(sys.modules, "pytesseract", None)  # makes "import pytesseract" fail
    ok, reason = ocr_module.ocr_available()
    assert not ok and "pip install pytesseract" in reason
    # Importing the module never fails, even without its optional dependencies.
    monkeypatch.setitem(sys.modules, "PIL", None)
    importlib.reload(ocr_module)
    assert ocr_module.ocr_available()[0] is False


def test_ocr_result_is_cached(monkeypatch, fresh_ocr_cache):
    calls = []
    monkeypatch.setattr(ocr_module, "_check", lambda: calls.append(1) or (True, "Tesseract 5 is installed."))
    assert ocr_module.ocr_available() == (True, "Tesseract 5 is installed.")
    assert ocr_module.ocr_available()[0]
    assert len(calls) == 1


@pytest.mark.parametrize(
    ("platform", "expected"),
    [("win32", "UB Mannheim"), ("darwin", "brew install tesseract"), ("linux", "apt install tesseract-ocr")],
)
def test_install_instructions_per_os(monkeypatch, platform, expected):
    monkeypatch.setattr(sys, "platform", platform)
    assert expected in ocr_module.install_instructions()


def test_windows_default_install_path_is_found(tmp_path, monkeypatch):
    pytesseract = pytest.importorskip("pytesseract")
    exe = tmp_path / "Programs" / "Tesseract-OCR" / "tesseract.exe"
    exe.parent.mkdir(parents=True)
    exe.write_bytes(b"")
    monkeypatch.setattr(sys, "platform", "win32")
    for var in ("ProgramFiles", "ProgramW6432", "ProgramFiles(x86)"):
        monkeypatch.setenv(var, str(tmp_path / "missing"))
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    monkeypatch.setattr(ocr_module.shutil, "which", lambda _cmd: None)
    monkeypatch.setattr(pytesseract.pytesseract, "tesseract_cmd", "tesseract")
    ocr_module._locate_tesseract(pytesseract)
    assert pytesseract.pytesseract.tesseract_cmd == str(exe)


# --------------------------------------------------------------------------- dispatcher


@pytest.mark.parametrize(
    ("name", "kind"),
    [
        ("book.epub", "epub"),
        ("notes.docx", "docx"),
        ("slides.pptx", "pptx"),
        ("page.html", "html"),
        ("notes.md", "md"),
        ("notes.txt", "txt"),
    ],
)
def test_every_loader_through_load_document(tmp_path, name, kind):
    path = tmp_path / name
    if kind == "epub":
        build_epub(path)
    elif kind == "docx":
        build_docx(path)
    elif kind == "pptx":
        build_pptx(path)
    elif kind == "html":
        path.write_text(PAGE, encoding="utf-8")
    elif kind == "md":
        path.write_text(MARKDOWN, encoding="utf-8")
    else:
        path.write_text(PLAIN, encoding="utf-8")
    progress_calls: list[tuple[float, str]] = []
    doc = load_document(path, source_name=f"Original {name}", progress=lambda f, m: progress_calls.append((f, m)))
    assert_sane(doc)
    assert doc.source_type == kind
    assert doc.source_name == f"Original {name}"
    assert len(all_blocks(doc)) >= 4
    assert progress_calls and all(0.0 <= f <= 1.0 for f, _ in progress_calls)
    assert load_document(path).id == doc.id  # same bytes, same id
    for block in all_blocks(doc):
        if block.kind not in (K.CODE, K.EQUATION):
            assert not _MD_MARKUP.search(block.text), block.text


@requires_tesseract
def test_image_through_load_document(tmp_path, fresh_ocr_cache):
    path = tmp_path / "upload.tmp"
    path.write_bytes(_image_bytes(render_text_image(["Photosynthesis makes sugar."]), "JPEG"))
    doc = load_document(path, source_name="Photo of board.jpg")
    assert_sane(doc)
    assert doc.source_type == "image" and doc.title == "Photo of board"
    assert "Photosynthesis makes sugar" in " ".join(texts(doc))


# --------------------------------------------------------------------------- shared text cleanup


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        # ordinary words broken by hyphenation are joined
        ("infor-\nmation", "information"),
        ("the move-\nment of ions", "the movement of ions"),
        ("multi-\nple", "multiple"),
        ("anti-\nbodies", "antibodies"),
        ("pre-\nvious", "previous"),
        ("ex-\nample", "example"),
        ("co-\nefficient", "coefficient"),
        ("self-\nish", "selfish"),
        ("high-\nlight", "highlight"),
        ("non-\nsense", "nonsense"),
        ("INFOR-\nMATION", "INFORMATION"),
        # compounds keep their hyphen
        ("self-\nesteem", "self-esteem"),
        ("Self-\nesteem", "Self-esteem"),
        ("well-\nknown", "well-known"),
        ("half-\nlife", "half-life"),
        ("cross-\nsection", "cross-section"),
        ("short-\nterm", "short-term"),
        ("non-\nzero", "non-zero"),
        ("pre-\nexisting", "pre-existing"),
        ("ex-\nwife", "ex-wife"),
        ("co-\nauthor", "co-author"),
        ("evidence-\nbased", "evidence-based"),
        ("state-of-the-\nart", "state-of-the-art"),
        ("interpreter-\nto-trace", "interpreter-to-trace"),
        ("multi-\nple-choice", "multiple-choice"),
        # a capital or a digit after the break: a real hyphen, and no stray space
        ("pre-\nColumbian", "pre-Columbian"),
        ("COVID-\n19", "COVID-19"),
        ("3-\nmethyl", "3-methyl"),
        # suspended hyphen, soft hyphen, dashes
        ("first-\nand second-order", "first- and second-order"),
        ("infor­\nmation", "information"),
        ("pages 10–\n20", "pages 10–20"),
        ("word—\nword", "word—word"),
        ("a dash -\nhere", "a dash - here"),
    ],
)
def test_line_break_hyphens(raw, expected):
    from athena.ingest.common import clean_display_text

    assert clean_display_text(raw) == expected


def test_line_break_hyphens_follow_the_documents_own_spelling():
    from athena.ingest.common import HyphenationEvidence, clean_display_text

    evidence = HyphenationEvidence(["Every co-op member votes.", "A cooperative store and a well-being plan."])
    assert clean_display_text("the co-\nop", evidence=evidence) == "the co-op"  # default would join
    assert clean_display_text("our co-\noperative", evidence=evidence) == "our cooperative"
    assert clean_display_text("well-\nbeing", evidence=evidence) == "well-being"
    evidence = HyphenationEvidence(["We call it TraceMonkey.", "TraceMonkey records traces."])
    assert clean_display_text("so Trace-\nMonkey starts", evidence=evidence) == "so TraceMonkey starts"
    evidence = HyphenationEvidence(["The selfesteem score."])  # (odd, but it is the book's spelling)
    assert clean_display_text("self-\nesteem", evidence=evidence) == "selfesteem"


def test_plain_text_line_break_hyphens_use_the_whole_text():
    blocks = plain_text_blocks(
        "Self-esteem grows slowly. Students with low self-\nesteem and the infor-\nmation they need.\n\n"
        "Every co-op meets on Friday; this co-\nop sells books."
    )
    assert [b.text for b in blocks] == [
        "Self-esteem grows slowly. Students with low self-esteem and the information they need.",
        "Every co-op meets on Friday; this co-op sells books.",
    ]


def test_control_characters_never_reach_the_text():
    from athena.ingest.common import clean_display_text, fix_characters

    # (math fonts in PDFs map big parentheses to control codes; NUL breaks voices)
    assert clean_display_text("log\x00 φ \x01 1−φ\x02") == "log φ 1−φ"
    assert clean_display_text("a\fb\x0bc\x1fd\x85e") == "a b c d e"
    assert fix_characters("def f():\n\tpass\x03") == "def f():\n\tpass"
    # Windows-1252 punctuation decoded as control codes is repaired, not dropped
    assert clean_display_text("FAIRY STORIES\x97MODERN TALES \x93quoted\x94 \x80 5\x81") == (
        "FAIRY STORIES\u2014MODERN TALES \u201cquoted\u201d \u20ac 5"
    )


@pytest.mark.parametrize("filler", [" ", "\xa0", "\t", "　", " \n"])
def test_long_runs_of_whitespace_are_cleaned_in_linear_time(filler):
    import time

    from athena.ingest.common import clean_display_text

    text = "start" + filler * 100_000 + "end"
    started = time.perf_counter()
    assert clean_display_text(text) == "start end"
    assert time.perf_counter() - started < 1.0  # was quadratic: 15+ seconds for 100k spaces


def test_line_break_hyphen_rules_stay_linear_time():
    import time

    from athena.ingest.common import HyphenationEvidence, clean_display_text

    started = time.perf_counter()
    clean_display_text("a-\n" * 50_000)
    clean_display_text("x" * 200_000 + "-\nyz")
    HyphenationEvidence(["a" * 200_000, "self-esteem " * 10_000])
    assert time.perf_counter() - started < 2.0


def test_markdown_patterns_stay_linear_time(tmp_path):
    import time

    started = time.perf_counter()
    markdown_blocks("`" * 50_000)
    looks_like_markdown("word " * 40_000)  # one 200k-character line
    path = tmp_path / "one line.txt"
    path.write_text("word " * 60_000, encoding="utf-8")
    load_document(path)
    assert time.perf_counter() - started < 3.0


def _write_epub2(path: Path, *, nav_points: str = "", guide: str = "") -> Path:
    """A hand-made EPUB 2 (contents page + one lab chapter), as InDesign-style exporters write them."""
    opf = (
        '<?xml version="1.0" encoding="UTF-8"?><package version="2.0" xmlns="http://www.idpf.org/2007/opf" '
        'unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Lab Manual</dc:title>'
        '<dc:identifier id="id">x</dc:identifier><dc:language>en</dc:language></metadata>'
        '<manifest><item id="toc" href="toc.xhtml" media-type="application/xhtml+xml"/>'
        '<item id="lab" href="lab.xhtml" media-type="application/xhtml+xml"/>'
        '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/></manifest>'
        f'<spine toc="ncx"><itemref idref="toc"/><itemref idref="lab"/></spine><guide>{guide}</guide></package>'
    )
    ncx = (
        '<?xml version="1.0"?><ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1"><head/>'
        f"<docTitle><text>Lab Manual</text></docTitle><navMap>{nav_points}</navMap></ncx>"
    )
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
            "</rootfiles></container>",
        )
        zf.writestr("OEBPS/content.opf", opf)
        zf.writestr("OEBPS/toc.ncx", ncx)
        zf.writestr("OEBPS/toc.xhtml", _xhtml("Contents", '<h1>Contents</h1><p><a href="lab.xhtml">Lab 1</a></p>'))
        zf.writestr("OEBPS/lab.xhtml", _xhtml("Lab", "<h1>Lab 1: Microscopes</h1><p>Clean the lens before use.</p>"))
    return path


_LAB_NAV_POINT = (
    '<navPoint id="n1" playOrder="1"><navLabel><text>Lab 1: Microscopes</text></navLabel>'
    '<content src="lab.xhtml"/></navPoint>'
)


def test_epub_guide_toc_label_on_the_only_chapter_is_not_trusted(tmp_path):
    # The guide calls both files the table of contents; only the page of links really is one.
    guide = '<reference type="toc" href="toc.xhtml" title="Contents"/><reference type="toc" href="lab.xhtml"/>'
    doc = load_document(_write_epub2(tmp_path / "guide.epub", nav_points=_LAB_NAV_POINT, guide=guide))
    assert_sane(doc)
    assert texts(doc) == ["Lab 1: Microscopes", "Clean the lens before use."]
    assert doc.warnings == []


def test_epub_with_an_empty_ncx_table_of_contents_opens(tmp_path):
    doc = load_document(_write_epub2(tmp_path / "empty-ncx.epub"))  # used to fail: "'Link' object is not iterable"
    assert_sane(doc)
    assert "Clean the lens before use." in texts(doc)


def test_docx_ruby_text_is_read_without_its_pronunciation_guide(tmp_path):
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    d = docx.Document()
    p = d.add_paragraph("Capital: ")
    p._p.append(
        parse_xml(
            f'<w:r {nsdecls("w")}><w:ruby><w:rubyPr><w:rubyAlign w:val="distributeSpace"/></w:rubyPr>'
            "<w:rt><w:r><w:t>とうきょう</w:t></w:r></w:rt><w:rubyBase><w:r><w:t>東京</w:t></w:r></w:rubyBase>"
            "</w:ruby></w:r>"
        )
    )
    p.add_run(".")
    path = tmp_path / "ruby.docx"
    d.save(str(path))
    doc = load_document(path)
    assert texts(doc) == ["Capital: 東京."]
    assert doc.warnings == []


@pytest.mark.parametrize(
    ("alt", "expected"),
    [
        ("Diagram of a plant cell", "Diagram of a plant cell"),
        ("IMG_2034.JPG", ""),
        ("K:\\Pics\\Chapter 3\\mitosis final.png", ""),
        ("/Users/ana/Desktop/cell.jpeg", ""),
        ("https://example.edu/images/cell.gif", ""),
        ("A close up of a logo Description generated with very high confidence", "A close up of a logo"),
        ("A diagram of a cell. Description automatically generated", "A diagram of a cell"),
    ],
)
def test_picture_descriptions_that_are_file_names_are_not_read(tmp_path, alt, expected):
    import docx
    import pptx
    from docx.shared import Inches
    from pptx.util import Inches as PptxInches

    from athena.ingest.common import clean_alt_text

    assert clean_alt_text(alt) == expected
    figures = [expected] if expected else []
    # Word
    d = docx.Document()
    d.add_paragraph("Look at the picture.")
    d.add_picture(io.BytesIO(png_bytes()), width=Inches(1))
    d.paragraphs[-1]._p.xpath(".//wp:docPr")[0].set("descr", alt)
    d.save(str(tmp_path / "pic.docx"))
    assert texts(load_document(tmp_path / "pic.docx"), K.FIGURE) == figures
    # PowerPoint (python-pptx itself writes the file name, as older PowerPoint versions did)
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Mitosis"
    picture = slide.shapes.add_picture(io.BytesIO(png_bytes()), PptxInches(1), PptxInches(2))
    picture._element.xpath(".//p:cNvPr")[0].set("descr", alt)
    prs.save(str(tmp_path / "pic.pptx"))
    assert texts(load_document(tmp_path / "pic.pptx"), K.FIGURE) == figures
    # web page / EPUB chapter
    blocks, _ = html_to_blocks(f'<p>Cells divide.</p><figure><img src="a.png" alt="{alt}"></figure>')
    assert [b.text for b in blocks if b.kind == K.FIGURE] == figures


_HANDOUT = [
    "Lab 3: Osmosis and Diffusion",
    "",
    "Purpose: observe how water moves across a membrane.",
    "Materials: potato cores, salt water, distilled water, scale.",
    "",
    "Weigh each potato core and record the mass in your notebook.",
    "Place three cores in salt water for thirty minutes.",
]


@requires_tesseract
@pytest.mark.parametrize("angle", [0, 8, 90, 180])
def test_photo_taken_at_an_angle_or_sideways_is_read(tmp_path, fresh_ocr_cache, angle):
    from PIL import Image

    # Before: 8 degrees, 90 and 180 all came out as nothing or as gibberish.
    photo = render_text_image(_HANDOUT, size=34, width=1500).rotate(
        angle, expand=True, fillcolor="white", resample=Image.Resampling.BICUBIC
    )
    path = tmp_path / "handout.jpg"
    path.write_bytes(_image_bytes(photo.convert("L"), "JPEG"))
    text = " ".join(texts(load_document(path)))
    for line in _HANDOUT:
        if line:
            assert line in text, (angle, text)


def test_ocr_reading_quality_measure():
    good = "Purpose: observe how water moves across a membrane."
    upside_down = "‘saynulty AUIY) JOJ 19M eS Ul $9109 891} Bde} yOoq9e}0U INOA UI] Sse BY} P1ODE1"  # (real output)
    assert not ocr_module._poorly_read(good)
    assert ocr_module._poorly_read("") and ocr_module._poorly_read(upside_down)
    assert ocr_module._word_count(good) > ocr_module._word_count(upside_down)


@pytest.mark.parametrize(
    ("text", "heading"),
    [("# Lab steps\n\nMix it. Wait.", "Lab steps"), ("Intro.\n\n# Results\n\nDone.", "Results")],
)
def test_pasted_text_with_a_single_heading(text, heading):
    from athena.ingest import document_from_text

    doc = document_from_text(text, "Other")
    blocks = [b for s in doc.sections for b in s.blocks]
    assert any(b.kind == BlockKind.HEADING and b.text == heading for b in blocks)
    assert not any(b.text.startswith("#") for b in blocks)


def test_pasted_hashtags_are_not_headings():
    from athena.ingest import document_from_text

    doc = document_from_text("Our #1 goal is reading.\n\nThen #2.", "Notes")
    assert all(b.kind != BlockKind.HEADING for s in doc.sections for b in s.blocks)
