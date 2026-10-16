"""Hostile and damaged files: every import ends in a document or a friendly IngestError,
quickly and without exhausting memory. All files are generated in code."""

from __future__ import annotations

import io
import time
import zipfile
from pathlib import Path

import pymupdf
import pytest
from pdf_factory import PdfBuilder

from athena.ingest import IngestError, load_document
from athena.ingest import common as common_mod
from athena.ingest.html import parse_html
from athena.models import Document


def texts(doc: Document) -> list[str]:
    return [b.text for s in doc.sections for b in s.blocks]


# --------------------------------------------------------------------------- sample files


def make_docx(path: Path) -> Path:
    import docx

    d = docx.Document()
    d.add_heading("Cell Transport", 1)
    d.add_paragraph("Osmosis moves water across a membrane.")
    d.save(str(path))
    return path


def make_pptx(path: Path) -> Path:
    import pptx

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Cell Transport"
    slide.placeholders[1].text = "Osmosis moves water"
    prs.save(str(path))
    return path


def make_epub(path: Path) -> Path:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("robustness")
    book.set_title("Cell Transport")
    book.set_language("en")
    chapter = epub.EpubHtml(title="Chapter 1", file_name="c1.xhtml")
    chapter.content = "<html><body><h1>Chapter 1</h1><p>Osmosis moves water across a membrane.</p></body></html>"
    book.add_item(chapter)
    book.toc = [epub.Link("c1.xhtml", "Chapter 1", "c1")]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter]
    epub.write_epub(str(path), book)
    return path


def make_pdf(path: Path) -> Path:
    b = PdfBuilder()
    b.heading("Cell Transport")
    b.paragraph("Osmosis moves water across a membrane. " * 8)
    return b.save(path)


MAKERS = {".docx": make_docx, ".pptx": make_pptx, ".epub": make_epub, ".pdf": make_pdf}


def add_member(src: Path, dst: Path, name: str, size: int, *, pad_existing: bool = False) -> Path:
    """Copy a zip and add ``size`` bytes of a highly compressible member (or pad an existing one)."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if pad_existing and info.filename == name:
                data += b" " * size  # (whitespace after the root element is still valid XML)
            zout.writestr(info.filename, data)
        if not pad_existing:
            zout.writestr(name, b"\0" * size)
    return dst


# --------------------------------------------------------------------------- zip bombs


@pytest.mark.parametrize(
    ("suffix", "member", "pad_existing"),
    [
        (".epub", "EPUB/images/huge.png", False),
        (".epub", "EPUB/c1.xhtml", True),
        (".docx", "word/media/image9.png", False),
        (".docx", "word/document.xml", True),
        (".pptx", "ppt/media/image9.png", False),
        (".pptx", "ppt/slides/slide1.xml", True),
    ],
)
def test_zip_bombs_are_refused_before_unpacking(tmp_path, monkeypatch, suffix, member, pad_existing):
    # The real limits are 1 GB in all, 320 MB per part and 256 MB of XML; scaled down here
    # so the test stays fast. Before the check, a 2 MB EPUB holding 2 GB of zeros took 20 s
    # and 4 GB of memory.
    monkeypatch.setattr(common_mod, "MAX_UNPACKED_BYTES", 8 << 20)
    monkeypatch.setattr(common_mod, "MAX_PART_BYTES", 6 << 20)
    monkeypatch.setattr(common_mod, "MAX_MARKUP_BYTES", 4 << 20)
    source = MAKERS[suffix](tmp_path / f"source{suffix}")
    bomb = add_member(source, tmp_path / f"bomb{suffix}", member, 7 << 20, pad_existing=pad_existing)
    assert bomb.stat().st_size < 200_000
    with pytest.raises(IngestError, match="far more than a normal document"):
        load_document(bomb)
    # an ordinary file with a modest picture still opens
    fine = add_member(source, tmp_path / f"fine{suffix}", "extra/picture.png", 1 << 20)
    assert load_document(fine).sections


def test_zip_limits_are_generous_for_real_files():
    assert common_mod.MAX_PART_BYTES >= 300 << 20  # a lecture video up to the 300 MB upload limit
    assert common_mod.MAX_UNPACKED_BYTES >= 1 << 30
    assert common_mod.MAX_MARKUP_BYTES >= 128 << 20


# --------------------------------------------------------------------------- damaged files


@pytest.mark.parametrize("suffix", [".docx", ".pptx", ".epub", ".pdf"])
@pytest.mark.parametrize("damage", ["empty", "garbage", "truncated", "half", "flipped"])
def test_damaged_files_give_a_friendly_error_or_a_document(tmp_path, suffix, damage):
    good = MAKERS[suffix](tmp_path / f"good{suffix}").read_bytes()
    bad = tmp_path / f"bad{suffix}"
    if damage == "empty":
        bad.write_bytes(b"")
    elif damage == "garbage":
        bad.write_bytes(b"\x00\x01garbage" * 100)
    elif damage == "truncated":
        bad.write_bytes(good[: int(len(good) * 0.97)])
    elif damage == "half":
        bad.write_bytes(good[: len(good) // 2])
    else:
        data = bytearray(good)
        for i in range(len(data) // 3, len(data) // 3 + 300):
            data[i] ^= 0x5A
        bad.write_bytes(bytes(data))
    started = time.perf_counter()
    try:
        doc = load_document(bad)
    except IngestError as exc:
        message = str(exc)
        assert message and "Traceback" not in message and len(message) < 400
    else:
        assert doc.sections and doc.title
    assert time.perf_counter() - started < 10


# --------------------------------------------------------------------------- mislabeled files


@pytest.mark.parametrize(
    ("real", "named", "kind"),
    [
        (".docx", "notes.pdf", "docx"),
        (".pdf", "notes.docx", "pdf"),
        (".pdf", "notes.txt", "pdf"),
        (".pptx", "slides.docx", "pptx"),
        (".epub", "book.docx", "epub"),
        (".docx", "notes.pptx", "docx"),
    ],
)
def test_file_with_the_wrong_extension_is_read_as_what_it_really_is(tmp_path, real, named, kind):
    path = MAKERS[real](tmp_path / f"real{real}")
    doc = load_document(path, source_name=named)
    assert doc.source_type == kind
    assert any("Osmosis moves water" in t for t in texts(doc))
    assert "but it is really" in doc.warnings[0]


def test_web_page_saved_as_pdf_gets_a_clear_message(tmp_path):
    # (what a student gets when a course site's download link returns its sign-in page)
    path = tmp_path / "lecture.pdf"
    path.write_text("<!DOCTYPE html><html><body><h1>Sign in</h1></body></html>", encoding="utf-8")
    with pytest.raises(IngestError, match="web page"):
        load_document(path)


def test_old_office_file_with_a_new_name_gets_a_clear_message(tmp_path):
    path = tmp_path / "chapter.pdf"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 2000)
    with pytest.raises(IngestError, match=r"older Word or PowerPoint file"):
        load_document(path)


def test_text_that_mentions_pdf_stays_text(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Notes on file formats.\n\nA PDF starts with %PDF-1.7 in its first line.\n", encoding="utf-8")
    doc = load_document(path)
    assert doc.source_type == "txt" and not doc.warnings


# --------------------------------------------------------------------------- pathological markup and text


def test_deeply_nested_unclosed_lists_are_read_in_linear_time():
    # Unclosed <ul><li> nesting forces the flat fallback reader, which used to check every
    # ancestor of every text: 20,000 levels took 8 minutes.
    html = "<html><body>" + "<ul><li>item" * 5000 + "</body></html>"
    started = time.perf_counter()
    content = parse_html(html)
    assert time.perf_counter() - started < 5
    assert len(content.blocks) == 5000 and content.blocks[0].text == "item"


def test_flat_fallback_still_drops_hidden_content():
    html = (
        "<html><body><p>Before</p><div hidden>" + "<div>" * 3000 + "SECRET" + "</div>" * 3000 + "</div>"
        "<h2>Heading</h2>" + "<div>" * 3000 + "deep <b>bold</b> text" + "</div>" * 3000 + "</body></html>"
    )
    content = parse_html(html)
    assert [(b.kind.value, b.text) for b in content.blocks] == [
        ("paragraph", "Before"),
        ("heading", "Heading"),
        ("paragraph", "deep bold text"),
    ]


def test_worksheet_answer_lines_do_not_slow_the_pdf_reader(tmp_path):
    # A packet of answer lines ("____") used to make the table-of-contents check quadratic.
    b = PdfBuilder()
    for page in range(8):
        if page:
            b.new_page()
        b.line(f"Question {page + 1}: explain what happens to a red blood cell in salt water.")
        for _ in range(44):
            b.line("_" * 95)
    path = b.save(tmp_path / "worksheet.pdf")
    started = time.perf_counter()
    load_document(path)
    assert time.perf_counter() - started < 3


@pytest.mark.parametrize(
    "text",
    [
        "مرحبا بالعالم. هذه فقرة عربية عن الخلايا.",  # right-to-left
        "שלום עולם. זו פסקה על תאים.",
        "细胞是生命的基本单位。细胞膜控制物质进出。",  # CJK
        "セル生物学の講義ノート。東京大学。",
        "Cells 🧬 divide 🔬 quickly 😀.",  # emoji
        "Café and naïve résumé.",  # combining marks
    ],
)
@pytest.mark.parametrize("suffix", [".txt", ".md", ".html", ".docx"])
def test_unusual_scripts_survive_import(tmp_path, text, suffix):
    import unicodedata

    path = tmp_path / f"notes{suffix}"
    if suffix == ".docx":
        import docx

        d = docx.Document()
        d.add_paragraph(text)
        d.save(str(path))
    elif suffix == ".html":
        path.write_text(f"<html><body><p>{text}</p></body></html>", encoding="utf-8")
    else:
        path.write_text(text + "\n", encoding="utf-8")
    doc = load_document(path)
    assert " ".join(texts(doc)) == unicodedata.normalize("NFC", text)


def test_cjk_pdf_text_survives(tmp_path):
    d = pymupdf.open()
    page = d.new_page()
    page.insert_text((72, 100), "细胞是生命的基本单位。", fontname="china-s", fontsize=14)
    page.insert_text((72, 130), "细胞膜控制物质进出细胞。", fontname="china-s", fontsize=14)
    path = tmp_path / "cjk.pdf"
    d.save(path)
    doc = load_document(path)
    assert "细胞是生命的基本单位" in "".join(texts(doc))


def test_pdf_without_any_text_or_pictures_explains_itself(tmp_path):
    d = pymupdf.open()
    for _ in range(3):
        d.new_page()
    path = tmp_path / "blank.pdf"
    d.save(path)
    doc = load_document(path)
    assert texts(doc) == []
    assert doc.warnings == ["No readable text was found in this PDF."]


def test_presentation_of_pictures_only_explains_itself(tmp_path):
    import pptx
    from pptx.util import Inches
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), "navy").save(buffer, "PNG")
    prs = pptx.Presentation()
    for _ in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        buffer.seek(0)
        slide.shapes.add_picture(buffer, Inches(1), Inches(1))
    path = tmp_path / "pictures.pptx"
    prs.save(str(path))
    doc = load_document(path)
    assert doc.warnings and "No readable text" in doc.warnings[0]
